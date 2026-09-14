"""标包与技术规范书解析模块。

从 ``包X_完整招标文件_<编码>.zip`` / ``包X_技术规范书_<编码>.zip`` 中按包提取：
项目名称、项目编号、分标名称、分标编号、包名称、包号、技术规范编码，
并把每个包内的技术规范书及关联文件归组输出。
"""
from __future__ import annotations

import json
import io
import logging
import re
import zipfile
from pathlib import Path

from sqlalchemy.orm import Session

from ..config import get_settings
from ..database import get_session_factory
from ..models import Task
from . import fallback_extractor
from .preprocess import ARCHIVE_EXTS, convert_to_markdown, extract_archive

logger = logging.getLogger(__name__)

PARSER_VERSION = 23
PACKAGE_ARCHIVE_PATTERN = re.compile(r"包\d+.*(?:完整招标文件|技术规范书)")
SCORE_SHEET_HINTS = ("技术评分标准", "技术详评细则", "评审细则", "详评细则")
SCORE_FILE_EXTS = {".xlsx", ".docx", ".doc"}
QUALIFICATION_SHEET_HINTS = ("专用资格", "资格后审", "资格预审")
QUALIFICATION_FILE_EXTS = {".xlsx", ".docx", ".doc"}
TECH_SPEC_NAME_HINTS = ("技术规范书", "技术规范", "技术规格", "技术参数")
TENDER_NAME_HINTS = ("招标公告", "招标文件", "投标文件格式")
NOTICE_NAME_HINTS = ("招标公告", "附件", "附表")
LIST_NAME_HINTS = ("货物清单", "物资清单", "清单", "汇总表")

META_PATTERNS: dict[str, list[str]] = {
    "project_name": [r"项目名称[：:]\s*([^：:\n\r]{2,120})", r"采购项目名称[：:]\s*([^：:\n\r]{2,120})"],
    "project_number": [r"项目编号[：:]\s*([A-Za-z0-9\-_—–]+)", r"采购项目编号[：:]\s*([A-Za-z0-9\-_—–]+)"],
    "bid_number": [r"招标编号[：:]\s*([A-Za-z0-9\-_—–]+)", r"招标项目编号[：:]\s*([A-Za-z0-9\-_—–]+)"],
    "subproject_name": [r"分标名称[：:]\s*([^：:\n\r]{2,120})", r"分标名称\s*[：:]\s*([^：:\n\r]{2,120})"],
    "subproject_number": [r"分标编号[：:]\s*([A-Za-z0-9\-_—–]+)"],
    "package_name": [r"包名称[：:]\s*([^：:\n\r]{2,40})"],
    "package_no": [r"包号[：:]\s*(包\d+)", r"包号[：:]\s*(\d+)\s*号?"],
}

META_LABELS = ("项目编号", "项目名称", "分标名称", "分标编号", "包名称", "包号", "技术规范编码")

META_CHINESE = {
    "project_name": "项目名称",
    "project_number": "项目编号",
    "bid_number": "招标编号",
    "subproject_name": "分标名称",
    "subproject_number": "分标编号",
    "package_name": "包名称",
    "package_no": "包号",
    "tech_spec_code": "技术规范编码",
}

ROW_FIELD_HEADERS = {
    "project_name": ("项目名称",),
    "project_number": ("项目编号",),
    "bid_number": ("招标编号",),
    "subproject_name": ("分标名称",),
    "subproject_number": ("分标编号",),
    "package_name": ("包名称",),
    "package_no": ("包号",),
    "tech_spec_code": ("技术规范编码", "技术规范书编码", "技术规范"),
}


def find_package_archives(*roots: Path) -> list[Path]:
    all_paths: list[Path] = []
    for root in roots:
        root = Path(root)
        if not root.exists():
            continue
        all_paths.extend(sorted(root.rglob("*")))
    files = [
        path
        for path in all_paths
        if path.is_file() and path.suffix.lower() in ARCHIVE_EXTS and PACKAGE_ARCHIVE_PATTERN.search(path.name)
    ]
    seen_stems: set[str] = set()
    unique_files: list[Path] = []
    for path in files:
        if path.stem in seen_stems:
            continue
        seen_stems.add(path.stem)
        unique_files.append(path)
    outer = [path for path in unique_files if "完整招标文件" in path.name]
    result = list(outer)
    if not outer:
        # 没有外层完整招标文件时，才单独解析“技术规范书”压缩包。
        result.extend(path for path in unique_files if "技术规范书" in path.name)
    for path in all_paths:
        if not (path.is_dir() and PACKAGE_ARCHIVE_PATTERN.search(path.name)):
            continue
        stem = path.name
        result = [item for item in result if not (item.name == stem or (item.is_file() and item.stem == stem))]
        result.append(path)
    return sorted(result, key=str)


def _ensure_full_extraction(work_dir: Path) -> Path:
    """从原始上传压缩包完整解压一份到 packages_raw，确保所有分标目录都被发现。"""
    input_dir = work_dir / "input"
    raw_dir = work_dir / "packages_raw"
    marker = raw_dir / ".extracted"
    if marker.exists():
        return raw_dir
    raw_dir.mkdir(parents=True, exist_ok=True)
    candidates = [
        path
        for path in sorted(input_dir.iterdir())
        if path.is_file() and path.suffix.lower() in ARCHIVE_EXTS and not PACKAGE_ARCHIVE_PATTERN.search(path.name)
    ]
    for archive in candidates:
        try:
            logger.info("[PACKAGE] 完整解压原始压缩包 %s -> %s", archive.name, raw_dir)
            extract_archive(archive, raw_dir, depth=0, warnings=[])
        except Exception as exc:
            logger.warning("完整解压失败 %s: %s", archive.name, exc)
    marker.write_text("done", encoding="utf-8")
    return raw_dir


def _load_combined_markdown(work_dir: Path) -> str | None:
    combined = work_dir / "combined.md"
    if combined.exists():
        return combined.read_text(encoding="utf-8", errors="replace")
    md_dir = work_dir / "markdown"
    if md_dir.exists():
        parts = [path.read_text(encoding="utf-8", errors="replace") for path in sorted(md_dir.glob("*.md"))]
        return "\n\n".join(parts) if parts else None
    return None


def _extract_tech_spec_chapter(text: str) -> tuple[str | None, str | None]:
    matches = [match for match in re.finditer(r"^第五章\s*技术规范书\s*$", text, re.M)]
    if not matches:
        return None, None
    start = matches[-1].start()
    end_match = re.search(r"^第[一二三四五六七八九十\d]+章", text[start + 1 :], re.M)
    end = start + 1 + (end_match.start() if end_match else min(60000, len(text) - start - 1))
    chapter = text[start:end].strip()
    return (chapter, "技术规范书（第五章）") if chapter else (None, None)


def _extract_tech_spec_pdf(source_path: Path, text: str, fallback_dir: Path) -> Path | None:
    """从原始 PDF 中切出技术规范书章节对应的页面，保持原格式。"""
    if not source_path.exists() or source_path.suffix.lower() != ".pdf":
        return None
    start_match = re.search(r"^第五章\s*技术规范书\s*$", text, re.M)
    if not start_match:
        return None
    page_positions = [(match.start(), int(match.group(1))) for match in re.finditer(r"<!--\s*page:\s*(\d+)\s*-->", text)]
    if not page_positions:
        return None
    start_page = next((page for pos, page in reversed(page_positions) if pos <= start_match.start()), page_positions[0][1])
    end_match = re.search(r"^第[一二三四五六七八九十\d]+章", text[start_match.end() :], re.M)
    end_pos = start_match.end() + (end_match.start() if end_match else 0)
    end_page = next((page for pos, page in reversed(page_positions) if pos <= end_pos), start_page)
    if end_page <= start_page:
        end_page = start_page + 1
    chapter_last_page = max(start_page, end_page - 2)
    try:
        import fitz

        output = fallback_dir / "技术规范书.pdf"
        with fitz.open(source_path) as source_doc:
            chapter_doc = fitz.open()
            for page_index in range(max(0, start_page - 1), min(source_doc.page_count, chapter_last_page)):
                chapter_doc.insert_pdf(source_doc, from_page=page_index, to_page=page_index)
            if chapter_doc.page_count:
                chapter_doc.save(output)
            chapter_doc.close()
        return output if output.exists() and output.stat().st_size > 0 else None
    except Exception as exc:
        logger.warning("技术规范书章节 PDF 切片失败 %s: %s", source_path.name, exc)
        return None


def _parse_text_score_rules(text: str) -> list[dict[str, str]]:
    start = text.find("技术部分评审细则")
    if start < 0:
        start = text.find("技术评分细则")
    if start < 0:
        return []
    end = text.find("商务部分评审细则", start)
    if end < 0:
        end = min(start + 9000, len(text))
    block = text[start:end]
    numeral_re = re.compile(r"^[一二三四五六七八九十]{1,3}$")
    rows: list[dict[str, str]] = []
    factor = ""
    content: list[str] = []
    for line in block.splitlines():
        stripped = line.strip()
        if "评审细则" in stripped and "技术部分" not in stripped:
            if factor and content:
                rows.append({"项目": "技术部分", "评审要素": factor, "评审内容": " ".join(content).strip()})
            factor = ""
            content = []
            break
        if re.match(r"^第[一二三四五六七八九十\d]+章", stripped):
            if factor and content:
                rows.append({"项目": "技术部分", "评审要素": factor, "评审内容": " ".join(content).strip()})
            factor = ""
            content = []
            break
        if numeral_re.match(stripped):
            if factor and content:
                rows.append({"项目": "技术部分", "评审要素": factor, "评审内容": " ".join(content).strip()})
            factor = ""
            content = []
            continue
        if not stripped or stripped in {"序号", "评审项目", "分值", "评分细则", "合计"}:
            continue
        if re.match(r"^<!--|^\d{1,3}$", stripped):
            continue
        if not factor:
            if "（满分" not in stripped and "评审细则" not in stripped:
                factor = stripped
        else:
            content.append(stripped)
    if factor and content:
        rows.append({"项目": "技术部分", "评审要素": factor, "评审内容": " ".join(content).strip()})
    return [row for row in rows if row["评审内容"]]


def _build_single_document_packages(task_id: str, work_dir: Path) -> dict | None:
    text = _load_combined_markdown(work_dir)
    if not text:
        return None
    source_path = ""
    with get_session_factory()() as db:
        task = db.get(Task, task_id)
        manifest = (task.manifest or {}) if task else {}
    for entry in manifest.get("files", []):
        if entry.get("status") == "converted" and entry.get("path"):
            source_path = str(entry["path"])
            break
    chapter, spec_name = _extract_tech_spec_chapter(text)
    files: list[dict] = []
    md_dir = work_dir / "markdown"
    md_paths = sorted(md_dir.glob("*.md")) if md_dir.exists() else []
    if md_paths:
        files.append(
            {
                "name": md_paths[0].name.replace(".md", ""),
                "rel_path": md_paths[0].name,
                "size": md_paths[0].stat().st_size,
                "kind": "tender_doc",
                "md_path": str(md_paths[0]),
                "source_path": str(md_paths[0]),
                "content_preview": "",
                "has_tech_params": False,
            }
        )
    if chapter and spec_name:
        fallback_dir = work_dir / "fallback_md"
        fallback_dir.mkdir(parents=True, exist_ok=True)
        spec_path = fallback_dir / "技术规范书.md"
        spec_path.write_text(chapter, encoding="utf-8")
        chapter_pdf = _extract_tech_spec_pdf(Path(source_path), text, fallback_dir) if source_path else None
        files.append(
            {
                "name": spec_name,
                "rel_path": spec_path.name,
                "size": spec_path.stat().st_size,
                "kind": "tech_spec",
                "md_path": str(spec_path),
                "source_path": str(chapter_pdf) if chapter_pdf else (source_path or str(spec_path)),
                "content_preview": re.sub(r"\s+", " ", chapter[:200]),
                "has_tech_params": any(keyword in chapter for keyword in ("技术参数特性表", "项目需求值", "投标人保证值")),
                "tech_spec_code": "",
            }
        )
    if not files:
        return None
    package = {
        "package_no": "包1",
        "package_name": "包1",
        "subproject_dir": "招标文件",
        "source_archive": "招标文件",
        "tech_spec_code": "",
        "files": files,
        "技术评分规则": _parse_text_score_rules(text),
    }
    subproject = {
        "分标名称": "未分组",
        "分标编号": "",
        "项目名称": "",
        "项目编号": "",
        "source_dir": "招标文件",
        "packages": [package],
    }
    return {
        "task_id": task_id,
        "project": {"分标名称": "未分组"},
        "packages": [],
        "subprojects": [subproject],
        "warnings": ["单文件招标文件，已按内嵌技术规范书生成兜底标包"],
    }


def parse_task_packages(task_id: str) -> dict:
    settings = get_settings()
    factory = get_session_factory()
    with factory() as db:
        task = db.get(Task, task_id)
        if task is None:
            raise ValueError(f"task not found: {task_id}")
        work_dir = Path(task.work_dir)

    input_dir = work_dir / "input"
    raw_dir = _ensure_full_extraction(work_dir)
    roots = (input_dir, raw_dir)
    archives = find_package_archives(*roots)
    if not archives:
        result = _build_single_document_packages(task_id, work_dir)
        if result is None:
            result = {
                "task_id": task_id,
                "project": {},
                "packages": [],
                "warnings": ["未找到 包X_完整招标文件_<编码>.zip 或 包X_技术规范书_<编码>.zip 压缩包"],
            }
        _write_packages_json(work_dir, result)
        return result

    packages: list[dict] = []
    for archive_path in archives:
        try:
            package = parse_package_archive(archive_path, work_dir)
            package["subproject_dir"] = _top_level_subproject_dir(archive_path, roots)
            packages.append(package)
        except Exception as exc:
            logger.warning("标包解析失败 %s: %s", archive_path.name, exc)
            packages.append(
                {
                    "package_no": _package_no_from_name(archive_path.name),
                    "package_name": "",
                    "subproject_dir": _top_level_subproject_dir(archive_path, roots),
                    "source_archive": archive_path.name,
                    "tech_spec_code": "",
                    "warnings": [f"{type(exc).__name__}: {exc}"],
                    "files": [],
                }
            )

    packages.sort(key=lambda item: _package_index(item.get("package_no", "")))
    project = _merge_project_meta(packages)
    markdown_dir = work_dir / "markdown"
    list_rows: list[dict[str, str]] = []
    list_paths: list[Path] = []
    notice_rows: list[dict[str, str]] = []
    if markdown_dir.exists():
        notice_texts = [
            path.read_text(encoding="utf-8", errors="replace")
            for path in markdown_dir.glob("*.md")
            if any(hint in path.name for hint in NOTICE_NAME_HINTS)
        ]
        list_texts = [
            path.read_text(encoding="utf-8", errors="replace")
            for path in markdown_dir.glob("*.md")
            if any(hint in path.name for hint in LIST_NAME_HINTS)
        ]
        list_paths = [path for path in markdown_dir.glob("*.md") if any(hint in path.name for hint in LIST_NAME_HINTS)]
        list_rows = _parse_list_rows("\n".join(list_texts))
        notice_rows = _parse_table_rows("\n".join(notice_texts))
        tender_meta = _extract_tender_meta("\n".join(notice_texts))
        codes = _extract_list_codes("\n".join(list_texts))
        for key, value in tender_meta.items():
            if value and not project.get(key):
                project[key] = value
        if codes and not project.get("tech_spec_code"):
            project["tech_spec_code"] = codes[0]
            project["tech_spec_codes"] = codes

    for package in packages:
        list_rows.extend(package.get("list_rows") or [])

    for package in packages:
        if package.get("subproject_dir") in {"", input_dir.name}:
            inferred = _infer_subproject_dir(package, list_paths)
            package["subproject_dir"] = inferred or package.get("分标名称") or "未分组"
        queries = [
            package.get("subproject_dir", ""),
            package.get("包号", ""),
            package.get("包名称", ""),
            package.get("技术规范编码", ""),
            package.get("source_archive", ""),
        ]
        row = _match_list_row(list_rows, [query for query in queries if query])
        if row:
            _merge_row_fields(package, row)
        notice_row = _best_row(notice_rows, [query for query in queries if query])
        if notice_row:
            notice_fields = _row_fields(notice_row)
            for key, value in notice_fields.items():
                if not package.get(key) and value and not _is_placeholder(value):
                    package[key] = value

    packages = _merge_duplicate_packages(packages)
    packages.sort(key=lambda item: _package_index(item.get("package_no", "")))

    for package in packages:
        package.update({label: package.get(key, "") for key, label in META_CHINESE.items()})
        if package.get("subproject_dir") in {"", "未分组"} and package.get("分标名称"):
            package["subproject_dir"] = package["分标名称"]
    for package in packages:
        name = str(package.get("包名称") or "")
        if (
            _is_placeholder(name)
            or name in {"序号", "名称", "项目单位"}
            or len(name) > 40
            or name.startswith(("项目", "分标", "包名称"))
        ):
            package["包名称"] = str(package.get("包号") or "")
            package["package_name"] = package["包名称"]
    subprojects = _build_subprojects(packages)
    _attach_score_rules(work_dir, packages, subprojects)
    _attach_qualification_rules(work_dir, packages, subprojects)
    _attach_score_weights(work_dir, packages, subprojects)
    for subproject in subprojects:
        for field in ("项目名称", "项目编号", "分标名称", "分标编号"):
            if not project.get(field) and subproject.get(field):
                project[field] = subproject[field]
    for key, label in META_CHINESE.items():
        if not project.get(label):
            project[label] = project.get(key, "")
    for subproject in subprojects:
        for field in ("分标名称", "分标编号", "项目名称", "项目编号"):
            if not subproject.get(field) and project.get(field):
                subproject[field] = project[field]
    for package in packages:
        for key, label in META_CHINESE.items():
            if not package.get(label):
                package[label] = package.get(key) or project.get(key) or project.get(label) or ""
    for target in [project, *subprojects, *packages]:
        if not target.get("项目编号") and target.get("招标编号"):
            target["项目编号"] = target["招标编号"]
            if "project_number" in target:
                target["project_number"] = target["招标编号"]
    result = {
        "task_id": task_id,
        "project": project,
        "packages": packages,
        "subprojects": subprojects,
        "warnings": [],
        "parser_version": PARSER_VERSION,
    }
    _write_packages_json(work_dir, result)
    return result


def load_or_parse(task_id: str) -> dict:
    with get_session_factory()() as db:
        task = db.get(Task, task_id)
        if task is None:
            raise ValueError(f"task not found: {task_id}")
        packages_json = Path(task.work_dir) / "packages.json"
    if packages_json.exists():
        cached = json.loads(packages_json.read_text(encoding="utf-8"))
        if cached.get("parser_version") == PARSER_VERSION:
            return cached
    return parse_task_packages(task_id)


def parse_package_archive(archive_path: Path, task_work_dir: Path) -> dict:
    stem = archive_path.stem
    dest = task_work_dir / "packages_work" / stem
    md_root = task_work_dir / "packages_md" / stem
    warnings: list[str] = []
    if archive_path.is_dir():
        dest = archive_path
    else:
        extract_archive(archive_path, dest, depth=0, warnings=warnings)

    entries: list[dict] = []
    texts: list[str] = []
    tender_texts: list[str] = []
    list_texts: list[str] = []
    list_rows: list[dict[str, str]] = []
    for file_path in sorted(dest.rglob("*")):
        if not file_path.is_file():
            continue
        try:
            md_path, _ = convert_to_markdown(file_path, md_root / f"{file_path.relative_to(dest).as_posix().replace('/', '__')}.md")
            content = Path(md_path).read_text(encoding="utf-8", errors="replace")
            texts.append(content)
            if any(hint in file_path.name for hint in TENDER_NAME_HINTS):
                tender_texts.append(content)
            if any(hint in file_path.name for hint in LIST_NAME_HINTS):
                list_texts.append(content)
        except Exception as exc:
            content = ""
            warnings.append(f"解析失败 {file_path.name}: {type(exc).__name__}: {exc}")

        if file_path.suffix.lower() == ".xlsx" and any(hint in file_path.name for hint in LIST_NAME_HINTS):
            list_rows.extend(_parse_xlsx_rows(file_path))
        is_tech_spec = (
            file_path.suffix.lower() not in ARCHIVE_EXTS
            and any(hint in file_path.name for hint in TECH_SPEC_NAME_HINTS)
        )
        is_tender_doc = any(hint in file_path.name for hint in ("招标文件", "招标公告", "投标文件格式"))
        has_tech_params = any(
            keyword in content for keyword in ("技术特性参数表", "技术参数特性表", "技术特性参数", "技术参数表")
        )
        file_code_match = re.search(
            r"[（(]([A-Za-z0-9]+-\d{5,}-\d{3,})[）)]",
            f"{file_path.parent.name} {file_path.name}",
        )
        entries.append(
            {
                "name": file_path.name,
                "rel_path": file_path.relative_to(dest).as_posix(),
                "size": file_path.stat().st_size,
                "kind": "tech_spec" if is_tech_spec else "tender_doc" if is_tender_doc else "other",
                "md_path": str(md_path) if content else "",
                "source_path": str(file_path),
                "content_preview": content[:400].strip(),
                "has_tech_params": has_tech_params,
                "tech_spec_code": file_code_match.group(1) if file_code_match else "",
            }
        )

    combined_text = "\n".join(texts)
    meta = _extract_tender_meta("\n".join(tender_texts) or combined_text)
    codes = _extract_list_codes("\n".join(list_texts))
    file_codes: list[str] = []
    for file_path in dest.rglob("*"):
        match = re.search(
            r"[（(]([A-Za-z0-9]+-\d{5,}-\d{3,})[）)]",
            f"{file_path.parent.name} {file_path.name}",
        )
        if match:
            code = match.group(1)
            if code not in file_codes:
                file_codes.append(code)
    for code in file_codes:
        if code not in codes:
            codes.append(code)
    if codes:
        meta["tech_spec_code"] = codes[0]
        meta["tech_spec_codes"] = codes
    package_no = _package_no_from_name(archive_path.name)
    if not meta.get("package_no"):
        meta["package_no"] = package_no
    if not meta.get("package_name"):
        meta["package_name"] = package_no or archive_path.stem

    package = {
        "package_no": meta.get("package_no", package_no),
        "package_name": meta.get("package_name", ""),
        "subproject_name": meta.get("subproject_name", ""),
        "subproject_number": meta.get("subproject_number", ""),
        "tech_spec_code": meta.get("tech_spec_code", ""),
        "tech_spec_codes": meta.get("tech_spec_codes", []),
        "subproject_dir": archive_path.parent.name,
        "source_archive": archive_path.name,
        "warnings": warnings[:20],
        "files": entries,
        "list_rows": list_rows,
        "meta": {key: value for key, value in meta.items() if value},
    }
    package.update({label: meta.get(key, "") for key, label in META_CHINESE.items()})
    return package


def _parse_list_rows(text: str) -> list[dict[str, str]]:
    return _parse_table_rows(text)


def _parse_xlsx_rows(path: Path) -> list[dict[str, str]]:
    """直接读取货物清单 xlsx，避免 Markdown 转换丢列。"""
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(path, data_only=True)
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            return []
        header = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        rows: list[dict[str, str]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            record = {
                str(header[col] or "").strip(): str(row[col]).strip() if col < len(row) and row[col] is not None else ""
                for col in range(len(header))
            }
            if any(record.values()):
                rows.append(record)
        return rows
    except Exception:
        return []


def _parse_table_rows(text: str) -> list[dict[str, str]]:
    """解析 Markdown 表格为行字典。"""
    rows: list[dict[str, str]] = []
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line.startswith("|"):
            index += 1
            continue
        header_cells = _split_row(line)
        if not any(cell for cell in header_cells):
            index += 1
            continue
        index += 1
        if index < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[index].strip()):
            index += 1
        while index < len(lines) and lines[index].strip().startswith("|"):
            cells = _split_row(lines[index].strip())
            row = {
                header_cells[col] if col < len(header_cells) else f"列{col + 1}": cells[col] if col < len(cells) else ""
                for col in range(max(len(header_cells), len(cells)))
            }
            rows.append(row)
            index += 1
    return rows


def _match_list_row(rows: list[dict[str, str]], queries: list[str]) -> dict[str, str] | None:
    return _best_row(rows, queries)


def _best_row(rows: list[dict[str, str]], queries: list[str]) -> dict[str, str] | None:
    best: dict[str, str] | None = None
    best_score = 0
    for row in rows:
        row_text = " ".join(value for value in row.values() if value)
        if not row_text:
            continue
        score = sum(1 for query in queries if query and query != "未分组" and query in row_text)
        if score > best_score:
            best = row
            best_score = score
    return best


def _merge_row_fields(package: dict, row: dict[str, str]) -> None:
    fields = _row_fields(row)
    for key, value in fields.items():
        if package.get(key):
            continue
        if value and not _is_placeholder(value):
            package[key] = value


def _row_fields(row: dict[str, str]) -> dict[str, str]:
    fields: dict[str, str] = {}
    for key, headers in ROW_FIELD_HEADERS.items():
        for cell_header, value in row.items():
            if any(header in cell_header for header in headers):
                clean = _clean_meta_value(value)
                if clean and _useful_meta_value(clean) and not _is_placeholder(clean):
                    fields[key] = clean
                    break
    return fields


def _is_placeholder(value: str) -> bool:
    if re.search(r"见投标人须知前附表|详见附表|详见招标文件|以招标文件为准|不适用|/|（如有）|待填|略", value):
        return True
    if "☑" in value or "补充" in value:
        return True
    if any(keyword in value for keyword in ("招标文件的全部内容", "愿以投标函", "按合同约定", "投标函", "全部内容")):
        return True
    if "项目单位" in value or "：" in value:
        return True
    if value.count("_") >= 2 or "XX" in value or "（" in value and "）" in value and len(value) <= 8:
        return True
    return False


def _infer_subproject_dir(package: dict, list_paths: list[Path]) -> str:
    """包压缩包位于根目录时，从货物清单文件名推断分标目录名。"""
    candidates: list[str] = []
    for path in list_paths:
        name = path.name
        match = re.search(r"货物清单_.*?_([\u4e00-\u9fffA-Za-z][\u4e00-\u9fffA-Za-z0-9\-]{1,40}?)\d{8}", name)
        if match:
            candidate = match.group(1)
            if candidate not in candidates:
                candidates.append(candidate)
        for token in re.findall(r"[\u4e00-\u9fff]{4,40}", name):
            if token not in ("货物清单", "招标文件", "招标公告") and token not in candidates:
                candidates.append(token)
    if not candidates:
        return ""
    queries = [str(package.get(key) or "") for key in ("包号", "包名称", "分标名称", "技术规范编码") if package.get(key)]
    for candidate in candidates:
        if any(query and candidate in query for query in queries):
            return candidate
    return candidates[0]


def _top_level_subproject_dir(archive_path: Path, roots: tuple[Path, ...]) -> str:
    """取上传根目录下第一层目录作为分标目录名。"""
    for root in roots:
        try:
            parts = archive_path.relative_to(root).parts
        except ValueError:
            continue
        if len(parts) >= 2:
            return parts[0]
    return archive_path.parent.name


def _merge_duplicate_packages(packages: list[dict]) -> list[dict]:
    """同一分标下 完整招标文件/技术规范书 属于同一包时合并为一个包。"""
    merged: dict[tuple[str, str], dict] = {}
    order: list[dict] = []
    for package in packages:
        key = (
            str(package.get("subproject_dir") or "未分组"),
            str(package.get("package_no") or _package_no_from_name(package.get("source_archive", "")) or ""),
        )
        existing = merged.get(key)
        if existing is None:
            merged[key] = package
            order.append(package)
            continue
        _merge_package(existing, package)
    return order


def _merge_package(target: dict, source: dict) -> None:
    known_files = {(file.get("rel_path"), file.get("name")) for file in target.get("files", [])}
    known_spec_codes = {
        file.get("tech_spec_code")
        for file in target.get("files", [])
        if file.get("kind") == "tech_spec" and file.get("tech_spec_code")
    }
    for file in source.get("files", []):
        if file.get("kind") == "tech_spec" and file.get("tech_spec_code") and file["tech_spec_code"] in known_spec_codes:
            continue
        if (file.get("rel_path"), file.get("name")) not in known_files:
            target.setdefault("files", []).append(file)
            known_files.add((file.get("rel_path"), file.get("name")))
            if file.get("kind") == "tech_spec" and file.get("tech_spec_code"):
                known_spec_codes.add(file["tech_spec_code"])
    for code in source.get("tech_spec_codes") or []:
        if code not in target.setdefault("tech_spec_codes", []):
            target["tech_spec_codes"].append(code)
    for warning in source.get("warnings") or []:
        if warning not in target.setdefault("warnings", []):
            target["warnings"].append(warning)
    source_archives = target.setdefault("source_archives", [])
    if target.get("source_archive") and target["source_archive"] not in source_archives:
        source_archives.append(target["source_archive"])
    if source.get("source_archive") and source["source_archive"] not in source_archives:
        source_archives.append(source["source_archive"])
    for key in ("包号", "包名称", "分标名称", "分标编号", "项目名称", "项目编号", "招标编号", "技术规范编码"):
        if not target.get(key) and source.get(key):
            target[key] = source[key]
    if not target.get("tech_spec_code") and source.get("tech_spec_code"):
        target["tech_spec_code"] = source["tech_spec_code"]


def _good_name(value: str) -> bool:
    value = str(value or "").strip()
    if len(value) < 2:
        return False
    if any(mark in value for mark in ("，", "、", "；", "：")):
        return False
    if value.startswith(("包名称", "项目名称", "分标名称", "投标人")):
        return False
    return True


def _package_match_name(package: dict) -> str:
    name = str(package.get("分标名称") or "")
    if not _good_name(name):
        name = str(package.get("subproject_dir") or "")
    return name


def _good_project_name(value: str) -> bool:
    value = str(value or "").strip()
    return len(value) >= 8 or ("-" in value and len(value) >= 6)


def _build_subprojects(packages: list[dict]) -> list[dict]:
    subprojects: dict[str, dict] = {}
    order: list[dict] = []
    for package in packages:
        subproject_dir = str(package.get("subproject_dir") or "未分组")
        project_name = _pkg_field(package, "项目名称", "project_name")
        subproject_number = _pkg_field(package, "分标编号", "subproject_number")
        if _good_project_name(project_name):
            identity = project_name
        elif _good_name(subproject_number):
            identity = subproject_number
        elif _pkg_field(package, "项目编号", "project_number"):
            identity = _pkg_field(package, "项目编号", "project_number")
        elif _pkg_field(package, "招标编号", "bid_number"):
            identity = _pkg_field(package, "招标编号", "bid_number")
        else:
            identity = str(package.get("source_archive") or "") or subproject_dir
        key = f"{subproject_dir}|{identity}"
        subproject = subprojects.get(key)
        if subproject is None:
            parsed_name = _pkg_field(package, "分标名称", "subproject_name")
            display_name = parsed_name if _good_name(parsed_name) else (subproject_dir if subproject_dir != "未分组" else identity)
            subproject = {
                "分标名称": display_name,
                "分标编号": _pkg_field(package, "分标编号", "subproject_number") or "",
                "项目名称": _pkg_field(package, "项目名称", "project_name") or "",
                "项目编号": _pkg_field(package, "项目编号", "project_number") or "",
                "source_dir": subproject_dir,
                "subproject_key": key,
                "packages": [],
            }
            subprojects[key] = subproject
            order.append(subproject)
        subproject["packages"].append(package)
        for field in ("分标名称", "分标编号", "项目名称", "项目编号"):
            if field == "分标名称":
                continue
            english = {
                "分标名称": "subproject_name",
                "分标编号": "subproject_number",
                "项目名称": "project_name",
                "项目编号": "project_number",
            }[field]
            if not subproject.get(field):
                subproject[field] = _pkg_field(package, field, english)
    return order


def _find_score_files(work_dir: Path) -> list[Path]:
    paths: list[Path] = []
    for root in (work_dir / "input", work_dir / "packages_raw"):
        if not root.exists():
            continue
        for path in sorted(root.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in SCORE_FILE_EXTS:
                continue
            if any(hint in path.name for hint in SCORE_SHEET_HINTS):
                paths.append(path)
    return paths


def _parse_score_mapping(workbook) -> dict[tuple[str, str], str]:
    if "技术评分标准目录" not in workbook.sheetnames:
        return {}
    sheet = workbook["技术评分标准目录"]
    mapping: dict[tuple[str, str], str] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name = str(row[0]).strip() if row and row[0] else ""
        package_no = str(row[1]).strip() if row and len(row) > 1 and row[1] else ""
        detail = str(row[2]).strip() if row and len(row) > 2 and row[2] else ""
        if name and detail:
            mapping[(name, package_no)] = detail
    return mapping


def _parse_score_detail_sheet(sheet) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    current_project = ""
    current_factor = ""
    for row in sheet.iter_rows(min_row=2, values_only=True):
        project = str(row[0]).strip() if row and row[0] else ""
        factor = str(row[1]).strip() if row and len(row) > 1 and row[1] else ""
        content = str(row[2]).strip() if row and len(row) > 2 and row[2] else ""
        if project:
            current_project = project
        if factor:
            current_factor = factor
        if project or factor or content:
            items.append({"项目": current_project, "评审要素": current_factor, "评审内容": content})
    return items


def _parse_score_docx(path: Path) -> tuple[list[dict[str, str]], str]:
    from docx import Document

    doc = Document(path)
    rows: list[dict[str, str]] = []
    for table in doc.tables:
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if not any("评审" in cell or "项目" in cell or "要素" in cell for cell in header):
            continue
        current_project = ""
        current_factor = ""
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            project = cells[0] if cells else ""
            factor = cells[1] if len(cells) > 1 else ""
            content = " ".join(cells[2:]) if len(cells) > 2 else ""
            if project:
                current_project = project
            if factor:
                current_factor = factor
            if project or factor or content:
                rows.append({"项目": current_project, "评审要素": current_factor, "评审内容": content})
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    return rows, text


def _parse_score_markdown(path: Path, work_dir: Path) -> tuple[list[dict[str, str]], str]:
    tmp_dir = work_dir / "score_tmp"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    md_path, _ = convert_to_markdown(path, tmp_dir / f"{path.stem}.md")
    text = md_path.read_text(encoding="utf-8", errors="replace")
    rows: list[dict[str, str]] = []
    for row in _parse_table_rows(text):
        def pick(*keys: str) -> str:
            for key in keys:
                for cell_header, value in row.items():
                    if key in cell_header and value:
                        return str(value).strip()
            return ""

        project = pick("项目")
        factor = pick("评审要素", "评审要点")
        content = pick("评审内容", "评审细则", "评分标准", "评分细则")
        if project or factor or content:
            rows.append({"项目": project, "评审要素": factor, "评审内容": content})
    return rows, text


def _extract_embedded_score_docs(docx_path: Path) -> list:
    """从招标文件 docx 中提取嵌入的 Word 对象（技术/商务评分细则）。"""
    from docx import Document

    docs: list = []
    try:
        with zipfile.ZipFile(docx_path) as archive:
            rels = archive.read("word/_rels/document.xml.rels").decode("utf-8", errors="replace")
            names = archive.namelist()
            for match in re.finditer(r'Id="(rId\d+)"[^>]*Type="[^"]*/oleObject[^"]*"[^>]*Target="([^"]+)"', rels):
                target = match.group(2)
                part = target if target.startswith("/") else f"word/{target}"
                part = part.lstrip("/")
                if part not in names:
                    continue
                try:
                    data = archive.read(part)
                except Exception:
                    continue
                index = data.find(b"PK\x03\x04")
                if index < 0:
                    continue
                try:
                    doc = Document(io.BytesIO(data[index:]))
                except Exception:
                    continue
                text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
                if "技术详评" in text or "评审要素" in text:
                    docs.append(doc)
    except Exception as exc:
        logger.warning("嵌入评分细则提取失败 %s: %s", docx_path.name, exc)
    return docs


def _parse_docx_score_table(table) -> list[dict[str, str]]:
    if not table.rows:
        return []
    header = [cell.text.strip() for cell in table.rows[0].cells]
    if not any("评审" in cell or "项目" in cell or "要素" in cell for cell in header):
        return []
    rows: list[dict[str, str]] = []
    current_project = ""
    current_factor = ""
    for row in table.rows[1:]:
        cells = [cell.text.strip() for cell in row.cells]
        project = cells[0] if cells else ""
        factor = cells[1] if len(cells) > 1 else ""
        content = " ".join(cells[2:]) if len(cells) > 2 else ""
        if project:
            current_project = project
        if factor:
            current_factor = factor
        if project or factor or content:
            rows.append({"项目": current_project, "评审要素": current_factor, "评审内容": content})
    return rows


def _parse_embedded_technical_score(docx_path: Path) -> dict[str, list[dict[str, str]]]:
    """解析嵌入 Word 中的技术详评细则：按“适用于XX分标”标题分组。"""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    result: dict[str, list[dict[str, str]]] = {}
    for doc in _extract_embedded_score_docs(docx_path):
        current_names: list[str] = []
        for child in doc.element.body.iterchildren():
            if child.tag == f"{{{word_ns}}}p":
                text = Paragraph(child, doc).text
                if "适用于" in text and "分标" in text:
                    segment = text.split("适用于", 1)[1].split("分标", 1)[0]
                    current_names = [name.strip() for name in re.split(r"[、，,]", segment) if name.strip()]
            elif child.tag == f"{{{word_ns}}}tbl":
                rows = _parse_docx_score_table(Table(child, doc))
                if rows:
                    for name in current_names:
                        result.setdefault(name, rows)
    return result


def _attach_score_rules(work_dir: Path, packages: list[dict], subprojects: list[dict]) -> None:
    """按 分标名称/分标编号 + 包号 精准匹配技术评分规则（支持 Excel 与 Word）。"""
    exact_map: dict[tuple[str, str], list[dict[str, str]]] = {}
    sheet_map: dict[str, list[dict[str, str]]] = {}
    word_candidates: list[tuple[list[dict[str, str]], str, str]] = []

    for path in _find_score_files(work_dir):
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            try:
                import openpyxl

                workbook = openpyxl.load_workbook(path, data_only=True)
            except Exception as exc:
                logger.warning("技术评分标准解析失败 %s: %s", path.name, exc)
                continue
            for (name, package_no), detail in _parse_score_mapping(workbook).items():
                if detail in workbook.sheetnames and (name, package_no) not in exact_map:
                    exact_map[(name, package_no)] = _parse_score_detail_sheet(workbook[detail])
            for sheet_name in workbook.sheetnames:
                if sheet_name == "技术评分标准目录" or sheet_name in sheet_map:
                    continue
                sheet_rows = _parse_score_detail_sheet(workbook[sheet_name])
                if sheet_rows:
                    sheet_map[sheet_name] = sheet_rows
        else:
            try:
                if suffix == ".docx":
                    rows, text = _parse_score_docx(path)
                else:
                    rows, text = _parse_score_markdown(path, work_dir)
            except Exception as exc:
                logger.warning("技术评分 Word 解析失败 %s: %s", path.name, exc)
                continue
            if rows:
                word_candidates.append((rows, text, path.name))

    embedded_map: dict[str, list[dict[str, str]]] = {}
    for root in (work_dir / "input", work_dir / "packages_raw"):
        if not root.exists():
            continue
        for tender_doc in root.rglob("*.docx"):
            if not any(keyword in tender_doc.name for keyword in ("招标文件", "招标公告")):
                continue
            try:
                embedded_map.update(_parse_embedded_technical_score(tender_doc))
            except Exception as exc:
                logger.warning("嵌入评分细则解析失败 %s: %s", tender_doc.name, exc)

    for package in packages:
        name = _package_match_name(package)
        number = str(package.get("分标编号") or "")
        package_no = str(package.get("包号") or "")
        rows = exact_map.get((name, package_no))
        if rows is None:
            for fallback_key in (("所有分标", "所有分包"), ("所有分标", "所有包"), ("全部", "全部")):
                if fallback_key in exact_map:
                    rows = exact_map[fallback_key]
                    break
        if rows is None:
            for key, candidate in sheet_map.items():
                if key == name or (name and name in key) or (number and number in key) or key in ("所有分标", "全部"):
                    rows = candidate
                    break
        if rows is None and len(sheet_map) == 1:
            rows = next(iter(sheet_map.values()))
        if rows is None:
            best_rows = None
            best_score = 0
            for candidate_rows, text, file_name in word_candidates:
                score = 0
                if name and name in text:
                    score += 3
                if number and number in text:
                    score += 3
                if package_no and package_no in text:
                    score += 1
                if name and name in file_name:
                    score += 2
                if number and number in file_name:
                    score += 2
                if score > best_score:
                    best_rows = candidate_rows
                    best_score = score
            rows = best_rows
        if rows is None:
            rows = embedded_map.get(name)
        if rows is None:
            for key, candidate in embedded_map.items():
                if key in name or name in key:
                    rows = candidate
                    break
        if rows:
            package["技术评分规则"] = rows

    for subproject in subprojects:
        subproject["技术评分规则"] = []
        seen: set[str] = set()
        for package in subproject.get("packages", []):
            for item in package.get("技术评分规则", []):
                key = json.dumps(item, ensure_ascii=False)
                if key not in seen:
                    seen.add(key)
                    subproject["技术评分规则"].append(item)


def _find_qualification_files(work_dir: Path) -> list[Path]:
    paths: list[Path] = []
    for root in (work_dir / "input", work_dir / "packages_raw"):
        if not root.exists():
            continue
        for path in sorted(root.rglob("*")):
            if not path.is_file() or path.suffix.lower() not in QUALIFICATION_FILE_EXTS:
                continue
            if any(hint in path.name for hint in QUALIFICATION_SHEET_HINTS):
                paths.append(path)
    return paths


def _parse_qualification_rows(rows: list[dict[str, str]]) -> dict[tuple[str, str], dict[str, str]]:
    result: dict[tuple[str, str], dict[str, str]] = {}
    for row in rows:
        name = str(row.get("分标名称") or row.get("标包名称") or "").strip()
        package_no = str(row.get("包号") or row.get("分包") or "").strip()
        if not name:
            continue
        fields: dict[str, str] = {"分标名称": name}
        if package_no:
            fields["包号"] = package_no
        for key, value in row.items():
            if key in ("分标名称", "包号", "分包"):
                continue
            clean = str(value).strip() if value else ""
            if clean:
                fields[key] = clean
        if fields:
            result[(name, package_no)] = fields
    return result


def _parse_qualification_xlsx(path: Path) -> dict[tuple[str, str], dict[str, str]]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True)
    for sheet in workbook.worksheets:
        header = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        if not any("分标名称" in cell for cell in header) or not any("包号" in cell or "分包" in cell for cell in header):
            continue
        rows: list[dict[str, str]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            record = {
                str(header[col] or "").strip(): str(row[col]).strip() if col < len(row) and row[col] is not None else ""
                for col in range(len(header))
            }
            rows.append(record)
        return _parse_qualification_rows(rows)
    return {}


def _parse_qualification_docx(path: Path) -> dict[tuple[str, str], dict[str, str]]:
    from docx import Document

    doc = Document(path)
    for table in doc.tables:
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if not any("分标名称" in cell for cell in header) or not any("包号" in cell or "分包" in cell for cell in header):
            continue
        rows: list[dict[str, str]] = []
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            record = {header[col]: cells[col] if col < len(cells) else "" for col in range(len(header))}
            rows.append(record)
        return _parse_qualification_rows(rows)
    return {}


def _parse_qualification_markdown(path: Path, work_dir: Path) -> dict[tuple[str, str], dict[str, str]]:
    tmp_dir = work_dir / "qualification_tmp"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    md_path, _ = convert_to_markdown(path, tmp_dir / f"{path.stem}.md")
    text = md_path.read_text(encoding="utf-8", errors="replace")
    rows: list[dict[str, str]] = []
    for row in _parse_table_rows(text):
        if any("分标名称" in key for key in row) and any("包号" in key or "分包" in key for key in row):
            rows.append(row)
    return _parse_qualification_rows(rows)


def _attach_qualification_rules(work_dir: Path, packages: list[dict], subprojects: list[dict]) -> None:
    """按 分标名称 + 包号 提取专用资格要求（支持 Excel 与 Word）。"""
    mapping: dict[tuple[str, str], dict[str, str]] = {}
    for path in _find_qualification_files(work_dir):
        try:
            if path.suffix.lower() == ".xlsx":
                parsed = _parse_qualification_xlsx(path)
            elif path.suffix.lower() == ".docx":
                parsed = _parse_qualification_docx(path)
            else:
                parsed = _parse_qualification_markdown(path, work_dir)
        except Exception as exc:
            logger.warning("专用资格要求解析失败 %s: %s", path.name, exc)
            continue
        for key, value in parsed.items():
            mapping.setdefault(key, value)
    for key, value in _qualification_rows_from_markdown(work_dir).items():
        mapping.setdefault(key, value)

    for package in packages:
        match_name = _package_match_name(package)
        key = (match_name, str(package.get("包号") or ""))
        fields = mapping.get(key)
        if fields is None:
            fields = mapping.get((match_name, ""))
        if fields is None:
            package_no = str(package.get("包号") or "")
            subproject_dir = str(package.get("subproject_dir") or "")
            for (name, candidate_no), candidate in mapping.items():
                if candidate_no in ("", package_no) and (name == subproject_dir or name == match_name or name in match_name):
                    fields = candidate
                    break
        if fields:
            package["专用资格要求"] = fields

    for subproject in subprojects:
        subproject["专用资格要求"] = None
        for package in subproject.get("packages", []):
            if package.get("专用资格要求"):
                subproject["专用资格要求"] = package["专用资格要求"]
                break


def _qualification_rows_from_markdown(work_dir: Path) -> dict[tuple[str, str], dict[str, str]]:
    rows: list[dict[str, str]] = []
    markdown_dir = work_dir / "markdown"
    if markdown_dir.exists():
        for md_path in markdown_dir.glob("*.md"):
            try:
                text = md_path.read_text(encoding="utf-8", errors="replace")
            except Exception:
                continue
            for row in _parse_table_rows(text):
                header_text = " ".join(row.keys())
                if "分标名称" not in header_text:
                    continue
                joined = " ".join(row.values())
                if any(keyword in header_text or keyword in joined for keyword in ("厂商类型", "业绩要求", "专用资格", "生产设备", "试验设备", "3C认证")):
                    rows.append(row)
    return _parse_qualification_rows(rows)


def _attach_score_weights(work_dir: Path, packages: list[dict], subprojects: list[dict]) -> None:
    """招标文件正文中的分值构成（技术/商务/价格权重）按分标名称提取。"""
    weights: dict[str, dict[str, str]] = {}
    markdown_dir = work_dir / "markdown"
    if markdown_dir.exists():
        for md_path in markdown_dir.glob("*.md"):
            try:
                text = md_path.read_text(encoding="utf-8", errors="replace")
            except Exception:
                continue
            for row in _parse_table_rows(text):
                name = str(row.get("分标名称") or "").strip()
                if not name:
                    continue
                tech = next((str(value).strip() for key, value in row.items() if "技术" in key and value), "")
                commercial = next((str(value).strip() for key, value in row.items() if "商务" in key and value), "")
                price = next((str(value).strip() for key, value in row.items() if "价格" in key and value), "")
                if tech or commercial or price:
                    weights[name] = {"技术": tech, "商务": commercial, "价格": price}

    for package in packages:
        name = _package_match_name(package)
        row = weights.get(name)
        if row is None:
            for key, candidate in weights.items():
                if key == str(package.get("subproject_dir") or "") or key in name:
                    row = candidate
                    break
        if row:
            package["评分权重"] = row
    for subproject in subprojects:
        subproject["评分权重"] = None
        for package in subproject.get("packages", []):
            if package.get("评分权重"):
                subproject["评分权重"] = package["评分权重"]
                break


def _pkg_field(package: dict, chinese: str, english: str) -> str:
    return str(package.get(chinese) or package.get(english) or "")


def _extract_tender_meta(text: str) -> dict:
    meta: dict[str, str] = {}
    for field, patterns in META_PATTERNS.items():
        if field == "tech_spec_code":
            continue
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                value = _clean_meta_value(match.group(1).strip())
                if value and _useful_meta_value(value) and not _is_placeholder(value):
                    meta[field] = value
                    break

    table_hints = {
        "subproject_number": "分标编号",
        "subproject_name": "分标名称",
        "package_no": "包号",
        "package_name": "包名称",
    }
    for field, hint in table_hints.items():
        if meta.get(field):
            continue
        values = _table_column_values(text, hint)
        if values:
            meta[field] = values[0]
    if meta.get("package_no"):
        if re.fullmatch(r"\d+", meta["package_no"]):
            meta["package_no"] = f"包{meta['package_no']}"
        elif not re.fullmatch(r"包\d+", meta["package_no"]):
            meta["package_no"] = ""
    for field in ("subproject_number", "project_number", "bid_number"):
        value = meta.get(field)
        if value and not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9\-_—–]{3,}", value):
            meta[field] = ""

    extract = fallback_extractor.extract_file(text, "招标文件")
    if not meta.get("project_name") and extract.get("project_name"):
        candidate = str(extract["project_name"]["primary_value"])
        if not _is_placeholder(candidate):
            meta["project_name"] = candidate
    if not meta.get("bid_number") and extract.get("bid_number"):
        candidate = str(extract["bid_number"]["primary_value"])
        if not _is_placeholder(candidate):
            meta["bid_number"] = candidate
    if not meta.get("project_number") and meta.get("bid_number"):
        meta["project_number"] = meta["bid_number"]
    return meta


def _extract_list_codes(text: str) -> list[str]:
    if not text:
        return []
    codes: list[str] = []
    for pattern in (
        r"技术规范编码[：:]\s*([A-Za-z0-9][A-Za-z0-9\-_]{5,})",
        r"技术规范(?:书)?编码\s*[：:]\s*([A-Za-z0-9][A-Za-z0-9\-_]{5,})",
        r"技术规范(?:书)?编码\s+([A-Za-z0-9][A-Za-z0-9\-_]{5,})",
    ):
        for match in re.finditer(pattern, text):
            code = match.group(1).strip()
            if code not in codes:
                codes.append(code)
    for value in _table_column_values(text, "技术规范"):
        if re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9\-_]{5,}", value) and value not in codes:
            codes.append(value)
    return codes


def _table_column_values(text: str, header_hint: str) -> list[str]:
    """从 Markdown 表格中按列名提取单元格值。"""
    lines = text.splitlines()
    for index, line in enumerate(lines):
        if not line.strip().startswith("|"):
            continue
        cells = _split_row(line)
        column = next((col for col, cell in enumerate(cells) if header_hint in cell), None)
        if column is None:
            continue
        values: list[str] = []
        for row_line in lines[index + 1 : index + 200]:
            if not row_line.strip().startswith("|"):
                break
            row_cells = _split_row(row_line)
            if len(row_cells) <= column:
                continue
            value = row_cells[column].strip()
            if value and _useful_meta_value(value) and not _is_placeholder(value) and value not in values:
                values.append(value[:120])
        if values:
            return values
    return []


def _split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _clean_meta_value(value: str) -> str:
    for label in META_LABELS:
        value = re.split(rf"(?={label}[：:])", value)[0]
        value = re.sub(rf"^{label}[：:]", "", value).strip()
    return value[:120]


def _useful_meta_value(value: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fffA-Za-z0-9]", value))


def _package_no_from_name(name: str) -> str:
    match = re.search(r"包(\d+)", name)
    return f"包{match.group(1)}" if match else ""


def _package_index(package_no: str) -> int:
    match = re.search(r"(\d+)", package_no or "")
    return int(match.group(1)) if match else 9999


def _merge_project_meta(packages: list[dict]) -> dict:
    project: dict[str, str] = {}
    for package in packages:
        meta = package.get("meta", {})
        for field in (
            "project_name",
            "project_number",
            "bid_number",
            "subproject_name",
            "subproject_number",
            "package_name",
            "package_no",
            "tech_spec_code",
        ):
            if not project.get(field) and meta.get(field):
                project[field] = meta[field]
    return project


def _write_packages_json(work_dir: Path, result: dict) -> None:
    (work_dir / "packages.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
