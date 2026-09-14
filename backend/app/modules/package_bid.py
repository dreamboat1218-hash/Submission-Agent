"""按包 + 按产品生成投标文件。

一个包生成一份 DOCX；用户勾选投哪几款产品后，每款产品生成：
技术偏差表、项目货物技术特性逐条逐项应答文件、项目货物组件材料配置表。
"""
from __future__ import annotations

import copy
import json
import logging
import posixpath
import re
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt
from docx.text.paragraph import Paragraph
from lxml import etree

from ..database import get_session_factory
from ..models import Task
from .preprocess import _decode_zip_name
from .bid_knowledge import build_package_knowledge, inject_knowledge_into_doc

logger = logging.getLogger(__name__)

TECH_PARAMS_HEADER = ["序号", "名称", "单位", "项目需求值或表述", "投标人保证值"]
COMPONENT_HEADER = ["序号", "组件材料名称", "规格型号", "单位", "数量", "备注"]
DEVIATION_HEADER = ["序号", "偏差事项", "招标文件要求", "投标文件响应", "偏差说明"]
SEQ_RE = re.compile(r"^\d+(?:\.\d+)*$")


def generate_package_bid(
    task_id: str,
    package_no: str,
    selected_paths: list[str],
    packages_data: dict | None = None,
    subproject: str | None = None,
) -> dict:
    factory = get_session_factory()
    with factory() as db:
        task = db.get(Task, task_id)
        if task is None:
            raise ValueError(f"task not found: {task_id}")
        work_dir = Path(task.work_dir)
        task_snapshot = {
            "extracted_data": task.extracted_data or {},
            "industry": task.industry or "IT",
        }

    if packages_data is None:
        packages_json = work_dir / "packages.json"
        packages_data = json.loads(packages_json.read_text(encoding="utf-8")) if packages_json.exists() else {}

    package = _find_package(packages_data, package_no, subproject)
    if package is None:
        raise ValueError(f"package not found: {package_no}" + (f"（分标：{subproject}）" if subproject else ""))

    project = _project_context_for_package(packages_data, package, subproject)
    tenderer_item = task_snapshot["extracted_data"].get("tenderer") if isinstance(task_snapshot["extracted_data"], dict) else None
    if isinstance(tenderer_item, dict) and tenderer_item.get("primary_value") and not project.get("招标人"):
        project["招标人"] = str(tenderer_item["primary_value"])
    files_by_md = {file.get("md_path", ""): file for file in package.get("files", [])}
    selected_files = [files_by_md[path] for path in selected_paths if path in files_by_md]
    if not selected_files:
        selected_files = [file for file in package.get("files", []) if file.get("kind") == "tech_spec" and file.get("has_tech_params")]
    if not selected_files:
        raise ValueError("未选择任何技术规范书")

    output_dir = work_dir / "output" / "packages" / _safe_name(package_no)
    output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = output_dir / f"投标文件_{_safe_name(package_no)}.docx"

    template_source = _package_template_source(work_dir, package)
    if not template_source:
        bid_format = task_snapshot["extracted_data"].get("bid_format", {})
        if isinstance(bid_format, dict) and isinstance(bid_format.get("primary_value"), dict):
            template_source = str(bid_format["primary_value"].get("source_path", "") or "")
    if template_source and Path(template_source).exists():
        try:
            return _generate_from_template(
                task_snapshot,
                template_source,
                package,
                selected_files,
                output_dir,
                docx_path,
                project,
                work_dir,
            )
        except Exception as exc:
            logger.warning("模版复制生成失败，降级为独立生成：%s", exc)

    doc = Document()
    _setup_document(doc)
    _add_cover(doc, project, package)
    doc.add_page_break()
    _add_toc(doc, selected_files)
    doc.add_page_break()

    _add_commercial(doc)
    doc.add_page_break()
    for index, file_entry in enumerate(selected_files, start=1):
        _add_product_response(doc, index, file_entry)
        if index < len(selected_files):
            doc.add_page_break()

    markdown = _render_markdown(project, package, selected_files)
    (output_dir / "投标文件.md").write_text(markdown, encoding="utf-8")
    doc.save(docx_path)

    return {
        "task_id": task_id,
        "package_no": package_no,
        "package_name": package.get("包名称") or package.get("package_name") or package_no,
        "bid_file_path": str(docx_path),
        "product_count": len(selected_files),
        "products": [file.get("name", "") for file in selected_files],
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }


def _generate_from_template(
    task_snapshot: dict,
    template_source: str,
    package: dict,
    selected_files: list[dict],
    output_dir: Path,
    docx_path: Path,
    project: dict | None = None,
    work_dir: Path | None = None,
) -> dict:
    """以附件投标格式.docx 为样式基座，插入招标文件技术部分并复制应答表。"""
    attachment = Path(__file__).resolve().parent.parent.parent / "templates" / "投标格式.docx"
    if not attachment.exists():
        raise RuntimeError("缺少标题样式附件：templates/投标格式.docx")

    doc = Document(attachment)
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    _clone_technical_from_source(doc, template_source)
    _add_technical_cover(doc, project or {}, package)
    _fill_table_header_values(doc, project or {}, package)
    _restyle_technical_heads(doc)

    blocks = (
        ("技术偏差表", ("项目货物技术特性逐条逐项应答文件",)),
        (
            "项目货物技术特性逐条逐项应答文件",
            (
                "组件材料选用说明",
                "投标人类别",
                "销售业绩",
                "检测(试验)报告",
                "认证证书",
                "许可证",
                "生产设备",
                "试验设备",
                "人力资源",
                "厂房环境",
                "其他资质要求",
            ),
        ),
    )
    for keyword, boundaries in blocks:
        start = _find_paragraph_element(doc, keyword)
        if start is None:
            continue
        block = _collect_block(start, boundaries)
        block_body = block[1:]
        if not block_body:
            continue
        type_index = blocks.index((keyword, boundaries)) + 1

        if keyword == "项目货物技术特性逐条逐项应答文件":
            _build_response_section(doc, start, block_body, selected_files, package)
        elif keyword == "技术偏差表":
            # 原文自带的偏差表标题与表格直接沿用，不再插入产品标题
            continue
        else:
            _duplicate_product_tables(doc, block_body, selected_files, type_index)

    _normalize_and_renumber_headings(doc)
    _restart_section_numbering(doc)
    _fill_response_tables(doc, selected_files, _load_goods_records(work_dir), project or {})
    knowledge_sections = 0
    knowledge_items = 0
    try:
        knowledge = build_package_knowledge(
            project or {},
            package,
            str(task_snapshot.get("industry") or ""),
        )
        knowledge_sections = inject_knowledge_into_doc(doc, knowledge)
        knowledge_items = sum(len(items) for items in knowledge.values())
    except Exception as exc:
        logger.warning("国网投标知识注入失败（不影响生成）：%s", exc)
    _drop_empty_heading_paragraphs(doc)
    doc.save(docx_path)
    _repair_cloned_images(docx_path, template_source)
    (output_dir / "投标文件.md").write_text(
        f"# 投标文件（{package.get('包名称') or package.get('package_no')}）\n\n"
        + "\n".join(f"## {file.get('name')}（{file.get('tech_spec_code', '')}）" for file in selected_files),
        encoding="utf-8",
    )
    return {
        "task_id": task_snapshot["extracted_data"].get("_meta", {}).get("task_id", ""),
        "package_no": package.get("包号") or package.get("package_no", ""),
        "package_name": package.get("包名称") or package.get("package_name") or "",
        "bid_file_path": str(docx_path),
        "product_count": len(selected_files),
        "products": [file.get("name", "") for file in selected_files],
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "template": "tender-format",
        "knowledge_sections": knowledge_sections,
        "knowledge_items": knowledge_items,
    }


def _repair_cloned_images(docx_path: Path, source_path: str) -> None:
    """补齐从招标文件技术部分克隆过来的图片，避免 docx 里出现悬空引用。"""
    try:
        with zipfile.ZipFile(docx_path) as target_zip:
            doc_xml = target_zip.read("word/document.xml")
            rels_xml = target_zip.read("word/_rels/document.xml.rels")
            target_media = {name for name in target_zip.namelist() if name.startswith("word/media/")}
        with zipfile.ZipFile(source_path) as source_zip:
            source_rels_xml = source_zip.read("word/_rels/document.xml.rels")
            source_media = {name: source_zip.read(name) for name in source_zip.namelist() if name.startswith("word/media/")}
    except (KeyError, FileNotFoundError):
        return

    nsmap = {"r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    source_root = etree.fromstring(source_rels_xml)
    source_map: dict[str, str] = {}
    for rel in source_root:
        if rel.get("Type", "").endswith("/image"):
            source_map[rel.get("Id")] = rel.get("Target")
    if not source_map:
        return

    doc_root = etree.fromstring(doc_xml)
    used_ids = set(doc_root.xpath("//@r:embed | //@r:link", namespaces=nsmap)) & set(source_map)
    if not used_ids:
        return

    rels_root = etree.fromstring(rels_xml)
    existing = {rel.get("Id"): (rel.get("Target"), rel.get("Type")) for rel in rels_root}
    rel_attrs = {rid: source_map[rid] for rid in used_ids if not source_map[rid].startswith(("http://", "https://", "/"))}
    if not rel_attrs:
        return

    replacements: dict[str, str] = {}
    to_add: dict[str, str] = {}
    next_id = 100
    for rid, target in sorted(rel_attrs.items()):
        reltype = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
        if rid in existing and existing[rid] == (target, reltype):
            continue
        if rid in existing:
            new_rid = f"rIdImage{next_id}"
            next_id += 1
            replacements[rid] = new_rid
            to_add[new_rid] = target
        else:
            to_add[rid] = target

    if replacements:
        for rid, new_rid in replacements.items():
            for element in doc_root.iter():
                for attr in (f"{{{nsmap['r']}}}embed", f"{{{nsmap['r']}}}link"):
                    if element.get(attr) == rid:
                        element.set(attr, new_rid)

    relationships = "http://schemas.openxmlformats.org/package/2006/relationships"
    for rid, target in sorted(to_add.items()):
        rel = etree.SubElement(rels_root, f"{{{relationships}}}Relationship")
        rel.set("Id", rid)
        rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")
        rel.set("Target", target)

    media_to_copy = {
        posixpath.normpath(posixpath.join("word", target.replace("\\", "/")))
        for target in rel_attrs.values()
    } - target_media
    if not to_add and not media_to_copy:
        return

    with zipfile.ZipFile(docx_path) as target_zip:
        items = target_zip.infolist()
        blobs = {item.filename: target_zip.read(item.filename) for item in items}
    content_types = _ensure_content_types(docx_path, media_to_copy)

    tmp_path = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as out:
        for item, blob in blobs.items():
            if item not in {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}:
                out.writestr(item, blob)
        out.writestr(
            "[Content_Types].xml",
            etree.tostring(content_types, xml_declaration=True, encoding="UTF-8", standalone=True),
        )
        out.writestr(
            "word/document.xml",
            etree.tostring(doc_root, xml_declaration=True, encoding="UTF-8", standalone=True),
        )
        out.writestr(
            "word/_rels/document.xml.rels",
            etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone=True),
        )
        for media_path in sorted(media_to_copy):
            blob = source_media.get(media_path)
            if blob is not None:
                out.writestr(media_path, blob)
    shutil.move(str(tmp_path), str(docx_path))


def _ensure_content_types(docx_path: Path, media_paths: set[str]) -> object:
    """确保包内包含新增图片所需的 jpeg/png 内容类型声明。"""
    with zipfile.ZipFile(docx_path) as target_zip:
        content_types_xml = target_zip.read("[Content_Types].xml")
    root = etree.fromstring(content_types_xml)
    content_types_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    existing = {
        element.get("Extension", "").lower()
        for element in root
        if element.tag == f"{{{content_types_ns}}}Default"
    }
    wanted: dict[str, str] = {}
    for media_path in media_paths:
        ext = posixpath.splitext(media_path)[1].lstrip(".").lower()
        if ext == "jpg":
            ext = "jpeg"
        if ext == "jpeg":
            wanted[ext] = "image/jpeg"
        elif ext == "png":
            wanted[ext] = "image/png"
        elif ext == "gif":
            wanted[ext] = "image/gif"
        elif ext == "bmp":
            wanted[ext] = "image/bmp"
    for ext, content_type in sorted(wanted.items()):
        if ext in existing:
            continue
        element = etree.Element(f"{{{content_types_ns}}}Default")
        element.set("Extension", ext)
        element.set("ContentType", content_type)
        root.insert(0, element)
    return root


def _clone_technical_from_source(doc: Document, source_path: str) -> None:
    source = Document(source_path)
    paragraphs = source.paragraphs
    start_index, start_marker = _find_technical_start(paragraphs)
    if start_index is None:
        raise RuntimeError("源模版中未找到技术部分")
    end_index = None
    if start_marker == "技术部分（单独成册）":
        end_index = next(
            (
                index
                for index in range(start_index + 1, len(paragraphs))
                if "投标保证保险（责任要求）" in paragraphs[index].text
            ),
            None,
        )
    start = paragraphs[start_index]._p
    end = paragraphs[end_index]._p if end_index is not None else None

    source_body = source.element.body
    elements: list = []
    capturing = False
    for element in source_body:
        if element is start:
            capturing = True
        if capturing:
            if end is not None and element is end:
                break
            elements.append(copy.deepcopy(element))

    target_body = doc.element.body
    sect_pr = target_body.find(qn("w:sectPr"))
    for element in elements:
        if sect_pr is not None:
            sect_pr.addprevious(element)
        else:
            target_body.append(element)


def _find_technical_start(paragraphs: list) -> tuple[int | None, str]:
    for marker in ("技术部分（单独成册）", "三、技术文件", "技术文件（格式）", "技术文件"):
        hits = [index for index, paragraph in enumerate(paragraphs) if marker in paragraph.text]
        if hits:
            return max(hits), marker
    return None, ""


def _package_template_source(work_dir: Path, package: dict) -> str:
    """优先使用当前分标自己的招标文件/投标文件格式作为投标模板源。"""
    files = package.get("files", []) or []

    def score(file_entry: dict) -> int:
        name = str(file_entry.get("name", ""))
        if "投标文件格式" in name:
            return 0
        if "招标文件" in name and "附件" not in name:
            return 1
        if "招标公告" in name:
            return 2
        return 3

    candidates = [file_entry for file_entry in files if file_entry.get("kind") == "tender_doc"]
    candidates.sort(key=score)
    for file_entry in candidates:
        source_path = str(file_entry.get("source_path", "") or "")
        if source_path and Path(source_path).exists() and Path(source_path).suffix.lower() == ".docx":
            return source_path
    for file_entry in candidates:
        source_path = str(file_entry.get("source_path", "") or "")
        if not source_path or not Path(source_path).exists() or Path(source_path).suffix.lower() != ".zip":
            continue
        try:
            with zipfile.ZipFile(source_path) as archive:
                for info in archive.infolist():
                    name = _decode_zip_name(info.filename)
                    if name.endswith(".docx") and ("招标文件" in name or "投标文件格式" in name):
                        tmp_dir = work_dir / "output" / "templates"
                        tmp_dir.mkdir(parents=True, exist_ok=True)
                        target = tmp_dir / f"{_safe_name(package.get('package_no') or 'pkg')}_format.docx"
                        target.write_bytes(archive.read(info))
                        return str(target)
        except Exception:
            continue
    return ""


def _find_paragraph_element(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if text in paragraph.text:
            return paragraph._p
    return None


def _collect_block(start_element, boundary_keywords: tuple[str, ...] = ()) -> list:
    body = start_element.getparent()
    elements: list = []
    capturing = False
    start_text = _element_text(start_element)
    start_major = _major_number(start_text)
    for element in body:
        if element is start_element:
            capturing = True
            elements.append(element)
            continue
        if not capturing:
            continue
        if element.tag == qn("w:p"):
            text = _element_text(element)
            if text and text != start_text:
                if any(keyword in text for keyword in boundary_keywords):
                    break
                number_match = re.match(r"^(\d+)[\.、]", text)
                if number_match and start_major is not None:
                    major = int(number_match.group(1))
                    if major != start_major:
                        break
        elements.append(element)
        if element.tag == qn("w:sectPr"):
            break
    return elements


def _major_number(text: str) -> int | None:
    match = re.match(r"^(\d+)[\.、]", text)
    return int(match.group(1)) if match else None


def _element_text(element) -> str:
    return "".join(node.text or "" for node in element.iter() if node.tag == qn("w:t"))


def _make_product_heading(doc: Document, product_index: int, file_entry: dict) -> object:
    paragraph = doc.add_paragraph()
    label = f"产品{product_index}：{file_entry.get('name', '')}"
    code = file_entry.get("tech_spec_code") or ""
    if code:
        label += f"（技术规范编码：{code}）"
    paragraph.add_run(label)
    try:
        paragraph.style = doc.styles["Heading 1"]
    except KeyError:
        pass
    _clear_run_formatting(paragraph)
    return paragraph._p


def _make_response_heading(doc: Document, package: dict, product_index: int, file_entry: dict) -> object:
    paragraph = doc.add_paragraph()
    package_no = package.get("包号") or package.get("package_no") or "包1"
    code = file_entry.get("tech_spec_code") or ""
    label = f"{package_no},技术规范编码：{code},项目货物技术特性逐条逐项应答文件"
    paragraph.add_run(label)
    try:
        paragraph.style = doc.styles["Heading 1"]
    except KeyError:
        pass
    _clear_run_formatting(paragraph)
    return paragraph._p


def _duplicate_product_tables(doc: Document, block_body: list, selected_files: list[dict], type_index: int) -> None:
    first_heading = _make_product_heading(doc, 1, selected_files[0])
    block_body[0].addprevious(first_heading)
    anchor = block_body[-1]
    for product_index, file_entry in enumerate(selected_files[1:], start=2):
        heading = _make_product_heading(doc, product_index, file_entry)
        copies = [copy.deepcopy(element) for element in block_body]
        first_copy = None
        for element in copies:
            anchor.addnext(element)
            if first_copy is None:
                first_copy = element
            anchor = element
        if first_copy is not None:
            first_copy.addprevious(heading)


def _build_response_section(
    doc: Document,
    start_element,
    block_body: list,
    selected_files: list[dict],
    package: dict,
) -> None:
    body = start_element.getparent()
    sub_indices = [
        index
        for index, element in enumerate(block_body)
        if element.tag == qn("w:p")
        and ("技术规范编码：XXXXXX" in _element_text(element) or re.match(r"^包\d+,技术规范编码：", _element_text(element)))
    ]
    if not sub_indices:
        _duplicate_product_tables(doc, block_body, selected_files, 2)
        return
    first_sub = sub_indices[0]
    pre = block_body[:first_sub]
    sub_tail = block_body[first_sub + 1 : (sub_indices[1] if len(sub_indices) > 1 else len(block_body))]
    for element in block_body[first_sub:]:
        body.remove(element)

    anchor = pre[-1] if pre else start_element
    for product_index, file_entry in enumerate(selected_files, start=1):
        heading = _make_response_heading(doc, package, product_index, file_entry)
        anchor.addnext(heading)
        anchor = heading
        for element in [copy.deepcopy(item) for item in sub_tail]:
            anchor.addnext(element)
            anchor = element


def _add_technical_cover(doc: Document, project: dict, package: dict) -> None:
    technical_el = _find_paragraph_element(doc, "技术部分（单独成册）")
    if technical_el is None:
        return
    package_no = package.get("包号") or package.get("package_no") or ""
    updates = [
        ("分标编号：", project.get("分标编号", "")),
        ("分标名称：", project.get("分标名称", "")),
        ("包", package_no),
        ("投标人：", "投标人：（盖章）"),
    ]
    for keyword, value in updates:
        paragraph = _find_cover_paragraph(doc, keyword)
        if paragraph is not None:
            original = paragraph.text.strip()
            if keyword == "包":
                prefix = f"{original.split('：')[0]}：" if "：" in original else "包号："
                label = f"{prefix}{value}"
            elif keyword == "投标人：":
                label = "投标人：（盖章）"
            else:
                label = f"{keyword}{value}"
            _set_paragraph_text_keep_format(paragraph, label)


_TABLE_HEADER_LABELS = ("项目名称", "项目编号", "分标名称", "分标编号", "包名称", "包号")


def _fill_table_header_values(doc: Document, project: dict, package: dict) -> None:
    """按当前分标/包填充表格头里的 项目名称/项目编号/分标名称/分标编号/包名称/包号。"""
    values = {
        "项目名称": str(project.get("项目名称") or ""),
        "项目编号": str(project.get("项目编号") or project.get("招标编号") or ""),
        "分标名称": str(project.get("分标名称") or ""),
        "分标编号": str(project.get("分标编号") or ""),
        "包名称": str(package.get("包名称") or project.get("包名称") or ""),
        "包号": str(package.get("包号") or package.get("package_no") or ""),
    }
    targets: list = []
    for paragraph in doc.paragraphs:
        if _header_has_labels(paragraph.text):
            targets.append(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if _header_has_labels(paragraph.text):
                        targets.append(paragraph)

    seen: set[int] = set()
    for paragraph in targets:
        if id(paragraph._p) in seen:
            continue
        seen.add(id(paragraph._p))
        pairs: list[tuple[str, str]] = []
        for label in _TABLE_HEADER_LABELS:
            if label in paragraph.text:
                pairs.append((f"{label}：", values.get(label) or ""))
        if pairs:
            _set_paragraph_text_with_pairs(paragraph, pairs)


def _header_has_labels(text: str) -> bool:
    return sum(1 for label in _TABLE_HEADER_LABELS if label in text) >= 3


def _set_paragraph_text_with_pairs(paragraph, pairs: list[tuple[str, str]]) -> None:
    """表格头：标签沿用原格式，填充值用宋体五号加下划线。"""
    template_rpr = None
    for run in list(paragraph.runs):
        if run._element.rPr is not None and template_rpr is None:
            template_rpr = copy.deepcopy(run._element.rPr)
        run._element.getparent().remove(run._element)

    for index, (label, value) in enumerate(pairs):
        label_run = paragraph.add_run(label)
        if template_rpr is not None:
            label_run._element.insert(0, copy.deepcopy(template_rpr))
        value_run = paragraph.add_run(str(value or ""))
        value_run.font.name = "宋体"
        value_run.font.size = Pt(10.5)
        value_run.underline = True
        rpr = value_run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:ascii"), "宋体")
        rfonts.set(qn("w:hAnsi"), "宋体")
        rfonts.set(qn("w:eastAsia"), "宋体")
        if index < len(pairs) - 1:
            space_run = paragraph.add_run(" ")
            if template_rpr is not None:
                space_run._element.insert(0, copy.deepcopy(template_rpr))


def _find_cover_paragraph(doc: Document, keyword: str):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if keyword == "包" and text.startswith("包") and "号" in text[:8]:
            return paragraph
        if text.startswith(keyword):
            return paragraph
    return None


def _set_paragraph_text_keep_format(paragraph, text: str) -> None:
    first_rpr = None
    for run in paragraph.runs:
        if run._element.rPr is not None and first_rpr is None:
            first_rpr = copy.deepcopy(run._element.rPr)
        run._element.getparent().remove(run._element)
    new_run = paragraph.add_run(text)
    if first_rpr is not None:
        new_run._element.insert(0, first_rpr)


def _make_formatted_paragraph(doc: Document, text: str, bold: bool = True, font: str = "黑体", size: int = 14) -> object:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = font
    run.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    return paragraph._p


def _set_heading_text_and_style(doc: Document, element, text: str, style_name: str) -> None:
    paragraph = Paragraph(element, doc)
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    paragraph.add_run(text)
    paragraph.style = doc.styles[style_name]
    _clear_run_formatting(paragraph)
    _strip_paragraph_auto_numbering(paragraph._p)


def _strip_paragraph_auto_numbering(element) -> None:
    p_pr = element.pPr
    if p_pr is None:
        return
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is not None:
        p_pr.remove(num_pr)


def _trim_to_technical(doc: Document) -> None:
    """只保留技术部分，移除封面、目录与商务文件。"""
    technical_el = _find_paragraph_element(doc, "技术部分（单独成册）")
    if technical_el is None:
        return
    body = doc.element.body
    for child in list(body):
        if child is technical_el:
            break
        body.remove(child)
    end_el = _find_paragraph_element(doc, "投标保证保险（责任要求）")
    if end_el is not None:
        removing = False
        for child in list(body):
            if child is end_el:
                removing = True
            if removing and child.tag != qn("w:sectPr"):
                body.remove(child)


def _drop_empty_heading_paragraphs(doc: Document) -> None:
    """删除空行却带标题样式的段落，避免出现空编号条目。"""
    for paragraph in list(doc.paragraphs):
        if paragraph.text.strip():
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if not style_name.startswith("Heading"):
            continue
        paragraph._p.getparent().remove(paragraph._p)


def _restyle_technical_heads(doc: Document) -> None:
    """按招标文件原文整理标题：括号/中文序号只加粗，其余标题统一为一级标题。"""
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        try:
            if "技术部分（单独成册）" in text:
                paragraph.style = doc.styles["Normal"]
                _clear_run_formatting(paragraph)
                continue
            if _has_manual_number_prefix(text):
                # （一）技术偏差表、三、技术文件 这类原文序号只保留加粗，不参与自动编号
                paragraph.style = doc.styles["Normal"]
                continue
            p_pr = paragraph._p.pPr
            outline = None
            if p_pr is not None and p_pr.outlineLvl is not None:
                try:
                    outline = int(p_pr.outlineLvl.val)
                except (TypeError, ValueError):
                    outline = None
            if outline is None:
                # 无大纲级别的加粗行属于表格标题，保持原样
                continue
            clean = re.sub(r"^\d+(?:\.\d+)*[\.、]?\s*", "", text)
            _set_heading_text_and_style(doc, paragraph._p, clean, "Heading 1")
        except KeyError:
            continue


def _has_manual_number_prefix(text: str) -> bool:
    """识别 （一）／三、 这类原文自带序号。"""
    return bool(
        re.match(r"^[（(][一二三四五六七八九十百\d]+[）)]", text)
        or re.match(r"^[一二三四五六七八九十]+、", text)
    )


def _looks_like_heading(text: str) -> bool:
    if len(text) > 60:
        return False
    if "。" in text or "；" in text or "，" in text:
        return False
    return True

def _clear_run_formatting(paragraph) -> None:
    """清除段落内直接字体格式，让标题完全套用附件样式。"""
    for run in paragraph.runs:
        r_pr = run._element.rPr
        if r_pr is not None:
            r_pr.getparent().remove(r_pr)


def _restart_section_numbering(doc: Document) -> None:
    """每个（一）（二）分节下的标题重新从 1 开始编号。"""
    try:
        numbering = doc.part.numbering_part.element
    except (AttributeError, NotImplementedError):
        return

    style_levels: dict[str, tuple[str, int]] = {}
    for level in range(1, 10):
        name = f"Heading {level}"
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        p_pr = style.element.pPr
        num_pr = p_pr.numPr if p_pr is not None else None
        if num_pr is None or num_pr.numId is None:
            continue
        base_id = str(num_pr.numId.val)
        abstract_id = None
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) != base_id:
                continue
            ref = num.find(qn("w:abstractNumId"))
            abstract_id = ref.get(qn("w:val")) if ref is not None else None
            break
        if abstract_id is None:
            continue
        ilvl = int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0
        style_levels[name] = (abstract_id, ilvl)
    if not style_levels:
        return

    used_ids = [
        int(num.get(qn("w:numId")))
        for num in numbering.findall(qn("w:num"))
        if str(num.get(qn("w:numId")) or "").isdigit()
    ]
    next_id = max(used_ids) + 1 if used_ids else 1
    section_ids: dict[str, int] = {}

    def _section_num_id(style_name: str) -> int:
        nonlocal next_id
        if style_name not in section_ids:
            abstract_id, ilvl = style_levels[style_name]
            num = OxmlElement("w:num")
            num.set(qn("w:numId"), str(next_id))
            ref = OxmlElement("w:abstractNumId")
            ref.set(qn("w:val"), abstract_id)
            num.append(ref)
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), str(ilvl))
            start_override = OxmlElement("w:startOverride")
            start_override.set(qn("w:val"), "1")
            override.append(start_override)
            num.append(override)
            numbering.append(num)
            section_ids[style_name] = next_id
            next_id += 1
        return section_ids[style_name]

    in_section = False
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if re.match(r"^[（(][一二三四五六七八九十百\d]+[）)]", text):
            section_ids.clear()
            in_section = True
            continue
        if not in_section:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        if style_name not in style_levels:
            continue
        ilvl = style_levels[style_name][1]
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = ilvl
        num_pr.get_or_add_numId().val = _section_num_id(style_name)


def _normalize_and_renumber_headings(doc: Document) -> None:
    """空行恢复正文，修正错误层级，并按 1/1.1/1.1.1 顺序统一编号。"""
    _fix_operation_manual(doc)
    _fix_sales_performance_list(doc)
    appendix_top = (
        "资质业绩文件",
        "投标人应用资质业绩凭证投标的说明",
    )
    prev_level = 1
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style else ""
        if not text:
            try:
                paragraph.style = doc.styles["Normal"]
            except KeyError:
                pass
            continue
        if text.startswith("（投标人若存在更名情况"):
            try:
                paragraph.style = doc.styles["Normal"]
            except KeyError:
                pass
            continue
        if text in {"中文核心及以上论文", "发明专利", "软著"}:
            try:
                paragraph.style = doc.styles["Normal"]
            except KeyError:
                pass
            continue
        if not style.startswith("Heading"):
            continue
        level = int(style.replace("Heading ", ""))
        if text.startswith("制造商授权书"):
            paragraph.style = doc.styles["Heading 2"]
            level = 2
        elif text.startswith(appendix_top):
            paragraph.style = doc.styles["Heading 1"]
            level = 1
        elif level >= 9:
            level = min(prev_level + 1, 4)
            try:
                paragraph.style = doc.styles[f"Heading {level}"]
            except KeyError:
                pass
        prev_level = level

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        style = paragraph.style.name if paragraph.style else ""
        if not style.startswith("Heading"):
            continue
        level = min(int(style.replace("Heading ", "")), 9)
        clean = re.sub(r"^\d+(?:\.\d+)*[\.、]?\s*", "", text)
        _set_heading_text_and_style(doc, paragraph._p, clean, f"Heading {level}")


OPERATION_STEPS = (
    ("1", "填报资质业绩信息"),
    ("2", "生成资质业绩凭证"),
    ("3", "下载资质业绩凭证"),
    ("4", "提交凭证单投标"),
)


def _fix_operation_manual(doc: Document) -> None:
    """操作手册步骤按示例还原为正文（1.-4.），不再作为一级标题。"""
    intro = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "二、操作介绍":
            intro = paragraph
            break
    if intro is None:
        return
    seen = False
    for paragraph in doc.paragraphs:
        if paragraph._p is intro._p:
            seen = True
            continue
        if not seen:
            continue
        clean = re.sub(r"^\d+(?:\.\d+)*[\.、]?\s*", "", paragraph.text.strip())
        for number, title in OPERATION_STEPS:
            if clean != title:
                continue
            try:
                paragraph.style = doc.styles["Normal"]
            except KeyError:
                pass
            for run in list(paragraph.runs):
                run._element.getparent().remove(run._element)
            run = paragraph.add_run(f"{number}.{title}")
            run.bold = True
            break

    for paragraph in doc.paragraphs:
        if paragraph.text.strip() in {"一、流程说明", "二、操作介绍"}:
            for run in paragraph.runs:
                run.bold = True


def _fix_sales_performance_list(doc: Document) -> None:
    """销售业绩的编制说明按示例改为 1./2. 有序列表。"""
    sales_heading = None
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("销售业绩（对应招标公告") or re.match(r"^\d+\.\s*销售业绩", text):
            sales_heading = paragraph
            break
    if sales_heading is None:
        return

    items: list = []
    in_section = False
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph._p is sales_heading._p:
            in_section = True
            continue
        if not in_section:
            continue
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name.startswith("Heading") and (text.startswith("检测") or re.match(r"^\d+\.\s*检测", text)):
            break
        if text != "编制说明：":
            continue
        following = doc.paragraphs[index + 1 : index + 3]
        if (
            len(following) == 2
            and following[0].text.strip().startswith("请据实填写")
            and following[1].text.strip().startswith("请按照")
        ):
            items = following
            break

    for number, paragraph in enumerate(items, start=1):
        _set_ordered_list_item(doc, paragraph, number)


def _set_ordered_list_item(doc: Document, paragraph, number: int) -> None:
    text = paragraph.text.strip()
    try:
        paragraph.style = doc.styles["List Paragraph"]
    except KeyError:
        pass
    paragraph_format = paragraph.paragraph_format
    paragraph_format.left_indent = Emu(777 * 635)
    paragraph_format.first_line_indent = Emu(-357 * 635)
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(f"{number}. {text}")
    run.bold = False


def _extract_tech_params(md_text: str) -> list[dict]:
    """从技术规范书 Markdown 提取技术参数特性表的序号/参数名称/项目需求值或表述。"""
    lines = md_text.splitlines()
    start = None
    end = len(lines)
    for index, line in enumerate(lines):
        if "技术参数特性表" in line and start is None:
            start = index + 1
            continue
        if start is not None and "组件材料配置表" in line:
            end = index
            break
    if start is None:
        return []
    region = lines[start:end]

    pipe_rows = [line for line in region if line.strip().startswith("|") and line.strip().endswith("|")]
    if pipe_rows:
        rows: list[dict] = []
        for line in pipe_rows:
            cells = _split_row(line)
            if len(cells) < 3:
                continue
            seq = cells[0].strip()
            if not seq or seq == "序号" or seq.startswith("---") or "填写" in seq:
                continue
            rows.append({"seq": seq, "name": cells[1].strip(), "value": cells[2].strip()})
        return rows

    non_empty = [line.strip() for line in region if line.strip()]
    header_index = None
    for index in range(len(non_empty) - 3):
        if non_empty[index] == "序号" and "参数名称" in non_empty[index + 1]:
            header_index = index
            break
    if header_index is None:
        return []

    rows = []
    index = header_index + 4
    while index < len(non_empty):
        seq = non_empty[index]
        if not SEQ_RE.match(seq):
            if seq.startswith("填写") or seq.startswith("规定"):
                break
            index += 1
            continue
        if index + 1 >= len(non_empty):
            break
        name = non_empty[index + 1]
        value_lines: list[str] = []
        index += 2
        while index < len(non_empty) and not SEQ_RE.match(non_empty[index]):
            if non_empty[index].startswith("填写") or non_empty[index].startswith("规定"):
                break
            value_lines.append(non_empty[index])
            index += 1
        if value_lines and value_lines[-1] == "—" and len(value_lines) > 1:
            value_lines.pop()
        rows.append({"seq": seq, "name": name, "value": "\n".join(value_lines)})
    return rows


def _load_goods_records(work_dir: Path | None) -> list[dict]:
    """从货物清单 xlsx 读取项目单位、项目名称等信息。"""
    if work_dir is None:
        return []
    try:
        candidates = sorted(Path(work_dir).glob("input/*货物清单*.xlsx"))
        if not candidates:
            return []
        import openpyxl

        workbook = openpyxl.load_workbook(str(candidates[0]), data_only=True)
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            return []
        header = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        records: list[dict] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            record = {
                str(header[col] or "").strip(): str(row[col]).strip() if col < len(row) and row[col] is not None else ""
                for col in range(len(header))
            }
            if record.get("项目名称") or record.get("项目单位"):
                records.append(record)
        return records
    except Exception:
        return []


def _match_goods_record(records: list[dict], file_entry: dict, project: dict) -> tuple[str, str]:
    default_unit = (
        records[0].get("项目单位", "")
        if records
        else (project.get("项目单位") or project.get("招标人") or project.get("tenderer") or "")
    )
    default_name = project.get("项目名称", "")
    haystack = " ".join(
        [
            str(file_entry.get("md_path", "")),
            str(file_entry.get("source_path", "")),
            str(file_entry.get("name", "")),
        ]
    )
    best = None
    best_length = 0
    for record in records:
        goods_name = record.get("项目名称", "")
        if not goods_name:
            continue
        if goods_name in haystack or haystack in goods_name:
            if len(goods_name) > best_length:
                best = record
                best_length = len(goods_name)
    if best is not None:
        return best.get("项目单位") or default_unit, best.get("项目名称") or default_name
    product_name = str(file_entry.get("name", ""))
    for record in records:
        material = record.get("物资名称", "")
        if material and material in product_name:
            return record.get("项目单位") or default_unit, record.get("项目名称") or default_name
    return default_unit, default_name


def _fill_response_tables(doc: Document, selected_files: list[dict], goods_records: list[dict], project: dict) -> None:
    """按产品把技术参数写入参数特性表，并补齐各表的项目单位/项目名称。"""
    params_tables: list = []
    component_tables: list = []
    env_tables: list = []
    for table in doc.tables:
        title = _table_title(table)
        if any(keyword in title for keyword in ("项目货物技术参数特性表", "技术特性参数表", "技术参数特性表")):
            params_tables.append(table)
        elif any(keyword in title for keyword in ("项目货物组件材料配置表", "组件材料配置表")):
            component_tables.append(table)
        elif "使用环境条件表" in title:
            env_tables.append(table)

    for index, file_entry in enumerate(selected_files):
        unit, name = _match_goods_record(goods_records, file_entry, project)
        if index < len(params_tables):
            _fill_project_row(params_tables[index], unit, name)
            _fill_param_table(params_tables[index], file_entry)
        if index < len(component_tables):
            _fill_project_row(component_tables[index], unit, name)
        if index < len(env_tables):
            _fill_project_row(env_tables[index], unit, name)
    if selected_files:
        unit, name = _match_goods_record(goods_records, selected_files[0], project)
        for table in doc.tables:
            _fill_any_project_row(table, unit, name)


def _table_title(table) -> str:
    rows = table.rows
    if not rows:
        return ""
    return _element_text(rows[0]._tr)


def _fill_project_row(table, unit: str, name: str) -> None:
    if len(table.rows) < 2:
        return
    row = table.rows[1]._tr
    for tc in row.findall(qn("w:tc")):
        text = _element_text(tc)
        if text.startswith("项目单位"):
            _set_cell_text(tc, f"项目单位：{unit}")
        elif text.startswith("项目名称"):
            _set_cell_text(tc, f"项目名称：{name}")


def _fill_any_project_row(table, unit: str, name: str) -> None:
    """补齐任意表格里独立的 项目单位/项目名称 单元格。"""
    for row in table.rows:
        for tc in row._tr.findall(qn("w:tc")):
            text = _element_text(tc)
            if "项目编号" in text or "分标名称" in text or "分标编号" in text or "包号" in text:
                continue
            if text.startswith("项目单位"):
                _set_cell_text(tc, f"项目单位：{unit}")
            elif text.startswith("项目名称"):
                _set_cell_text(tc, f"项目名称：{name}")


def _fill_param_table(table, file_entry: dict) -> None:
    md_path = file_entry.get("md_path", "")
    params = _extract_params_from_doc(file_entry)
    if not params and md_path and Path(md_path).exists():
        params = _extract_tech_params(Path(md_path).read_text(encoding="utf-8", errors="replace"))
    if not params:
        return

    table_element = table._tbl
    trs = table_element.findall(qn("w:tr"))
    header_index = None
    seq_col, name_col, unit_col, value_col, guarantee_col = 0, 1, 2, 3, 4
    for index, tr in enumerate(trs):
        cells = tr.findall(qn("w:tc"))
        texts = [_element_text(tc) for tc in cells]
        joined = " ".join(texts)
        if "序号" in joined and any(keyword in joined for keyword in ("项目需求值或表述", "招标人要求值", "要求值", "需求值")):
            header_index = index
            for cell_index, cell_text in enumerate(texts):
                if "序号" in cell_text:
                    seq_col = cell_index
                elif "保证值" in cell_text:
                    guarantee_col = cell_index
                elif "名称" in cell_text:
                    name_col = cell_index
                elif "单位" in cell_text:
                    unit_col = cell_index
                elif "要求值" in cell_text or "需求值" in cell_text:
                    value_col = cell_index
            break
    if header_index is None or header_index + 1 >= len(trs):
        return

    note_row = None
    if "编制说明" in _element_text(trs[-1]):
        note_row = copy.deepcopy(trs[-1])
    template_row = copy.deepcopy(trs[header_index + 1])
    for tr in trs[header_index + 1 :]:
        table_element.remove(tr)

    anchor = trs[header_index]
    for param in params:
        new_row = copy.deepcopy(template_row)
        cells = new_row.findall(qn("w:tc"))
        max_col = max(seq_col, name_col, unit_col, value_col, guarantee_col)
        if len(cells) > max_col:
            _set_cell_text(cells[seq_col], param.get("seq", ""))
            _set_cell_text(cells[name_col], param.get("name", ""))
            _set_cell_text(cells[unit_col], param.get("unit", ""))
            _set_cell_text(cells[value_col], param.get("value", ""))
            _set_cell_text(cells[guarantee_col], "")
        anchor.addnext(new_row)
        anchor = new_row
    if note_row is not None:
        anchor.addnext(note_row)


def _extract_params_from_doc(file_entry: dict) -> list[dict]:
    """直接从技术规范书原文档的表格提取参数，避免 Markdown 转换丢列。"""
    source = file_entry.get("source_path") or file_entry.get("md_path")
    if not source or not Path(source).exists():
        return []
    try:
        from docx import Document as DocxDocument

        source_path = Path(source)
        if source_path.suffix.lower() == ".doc":
            import subprocess
            import tempfile

            with tempfile.TemporaryDirectory() as temp_dir:
                result = subprocess.run(
                    [
                        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
                        "--headless",
                        "--convert-to",
                        "docx",
                        "--outdir",
                        temp_dir,
                        str(source_path),
                    ],
                    capture_output=True,
                    timeout=120,
                )
                candidates = sorted(Path(temp_dir).glob("*.docx"))
                if result.returncode != 0 or not candidates:
                    return []
                doc = DocxDocument(candidates[0])
        else:
            doc = DocxDocument(source_path)
    except Exception as exc:
        logger.warning("技术规范书表格参数提取失败 %s: %s", source, exc)
        return []
    return _params_from_docx_tables(doc)


def _params_from_docx_tables(doc) -> list[dict]:
    for table in doc.tables:
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if not (
            any("序号" in cell for cell in header)
            and any(("名称" in cell or "项目" in cell) for cell in header)
            and any(("要求值" in cell or "需求值" in cell) for cell in header)
        ):
            continue
        seq_col = next((index for index, cell in enumerate(header) if "序号" in cell), 0)
        name_col = next((index for index, cell in enumerate(header) if "名称" in cell or "项目" in cell), 1)
        value_col = next((index for index, cell in enumerate(header) if "要求值" in cell or "需求值" in cell), 3)
        unit_col = next((index for index, cell in enumerate(header) if "单位" in cell), None)
        params: list[dict] = []
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            seq = cells[seq_col] if len(cells) > seq_col else ""
            if not seq or seq == "序号":
                continue
            name = cells[name_col] if len(cells) > name_col else ""
            value = cells[value_col] if len(cells) > value_col else ""
            params.append(
                {
                    "seq": seq,
                    "name": name,
                    "unit": cells[unit_col] if unit_col is not None and len(cells) > unit_col else "",
                    "value": value,
                }
            )
        if params:
            return params
    return []


def _set_cell_text(tc, text: str) -> None:
    paragraphs = tc.findall(qn("w:p"))
    first = paragraphs[0] if paragraphs else None
    if first is None:
        first = etree.SubElement(tc, qn("w:p"))
    for paragraph in paragraphs[1:]:
        tc.remove(paragraph)

    template_rpr = None
    for run in first.findall(qn("w:r")):
        rpr = run.find(qn("w:rPr"))
        if rpr is not None and template_rpr is None:
            template_rpr = copy.deepcopy(rpr)
        first.remove(run)

    lines = str(text).split("\n")
    for line_index, line in enumerate(lines):
        if line_index > 0:
            run = etree.SubElement(first, qn("w:r"))
            etree.SubElement(run, qn("w:br"))
        run = etree.SubElement(first, qn("w:r"))
        if template_rpr is not None:
            run.append(copy.deepcopy(template_rpr))
        text_element = etree.SubElement(run, qn("w:t"))
        text_element.text = line
        text_element.set(qn("xml:space"), "preserve")


def _apply_attachment_styles(docx_path: Path) -> None:
    attachment = Path(__file__).resolve().parent.parent.parent / "templates" / "投标格式.docx"
    if not attachment.exists():
        return
    with zipfile.ZipFile(attachment) as source, zipfile.ZipFile(docx_path) as target:
        styles_xml = source.read("word/styles.xml")
    tmp_path = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as source_zip, zipfile.ZipFile(tmp_path, "w") as target_zip:
        for item in source_zip.infolist():
            data = source_zip.read(item.filename)
            if item.filename == "word/styles.xml":
                data = styles_xml
            target_zip.writestr(item, data)
    shutil.move(str(tmp_path), str(docx_path))


def _find_package(packages_data: dict, package_no: str, subproject: str | None = None) -> dict | None:
    candidates: list[dict] = []
    for package in packages_data.get("packages", []):
        if package not in candidates:
            candidates.append(package)
    for subproject_item in packages_data.get("subprojects", []):
        for package in subproject_item.get("packages", []):
            if package not in candidates:
                candidates.append(package)

    matches = [
        package
        for package in candidates
        if package.get("package_no") == package_no or package.get("包号") == package_no
    ]
    if subproject:
        for package in matches:
            if _package_matches_subproject(package, subproject):
                return package
    return matches[0] if matches else None


def _package_matches_subproject(package: dict, subproject: str) -> bool:
    subproject = str(subproject or "").strip()
    if not subproject:
        return True
    return any(
        str(package.get(key) or "") == subproject
        for key in ("subproject_dir", "subproject_key", "分标名称")
    ) or str(package.get("source_archive", "")).startswith(subproject)


def _project_context_for_package(packages_data: dict, package: dict, subproject: str | None = None) -> dict:
    """组装按分标区分的项目上下文，避免同名包号串用其他分标内容。"""
    project = dict(packages_data.get("project", {}) or {})
    if subproject:
        for subproject_item in packages_data.get("subprojects", []) or []:
            if not _package_matches_subproject(subproject_item, subproject):
                continue
            for key in ("项目名称", "项目编号", "招标编号", "分标名称", "分标编号", "包名称", "包号"):
                if subproject_item.get(key) and not _is_garbage_field(subproject_item.get(key)):
                    project[key] = subproject_item[key]
            break
    for key in ("项目名称", "项目编号", "招标编号", "分标名称", "分标编号", "包名称", "包号"):
        if package.get(key) and not _is_garbage_field(package.get(key)):
            project[key] = package[key]
    if _is_garbage_field(project.get("分标名称")):
        for subproject_item in packages_data.get("subprojects", []) or []:
            if not _package_matches_subproject(subproject_item, subproject):
                continue
            candidate = subproject_item.get("分标名称") or subproject_item.get("source_dir")
            if candidate and not _is_garbage_field(candidate):
                project["分标名称"] = candidate
            break
    return project


def _is_garbage_field(value) -> bool:
    text = str(value or "").strip()
    if not text:
        return True
    if any(mark in text for mark in ("，", "、", "；", "：")):
        return True
    if text.startswith(("包名称", "项目名称", "分标名称", "投标人")) or len(text) > 120:
        return True
    return False


def _safe_name(name: str) -> str:
    return "".join(ch for ch in str(name) if ch.isalnum() or ch in ("-", "_", " ", "包", "号")).strip() or "package"


def _setup_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _add_cover(doc: Document, project: dict, package: dict) -> None:
    for _ in range(5):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("投 标 文 件")
    run.font.name = "SimHei"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(34)
    doc.add_paragraph()
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    package_label = package.get("包名称") or package.get("package_name") or package.get("package_no", "")
    label = str(package_label or "")
    if (
        not label
        or len(label) > 40
        or "：" in label
        or label.startswith(("项目", "分标", "包名称"))
        or any(keyword in label for keyword in ("招标文件", "投标函", "全部内容"))
    ):
        label = str(package.get("包号") or package.get("package_no") or "")
    subtitle.add_run(f"（{label}）").font.size = Pt(16)
    for _ in range(4):
        doc.add_paragraph()
    for label, value in (
        ("项目名称", project.get("项目名称", "")),
        ("项目编号", project.get("项目编号", "")),
        ("分标名称", project.get("分标名称", "")),
        ("分标编号", project.get("分标编号", "")),
        ("包号", package.get("包号") or package.get("package_no", "")),
    ):
        if value:
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(f"{label}：{value}").font.size = Pt(14)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("投标人：____________（盖章）").font.size = Pt(14)


def _add_toc(doc: Document, selected_files: list[dict]) -> None:
    heading = doc.add_paragraph("目 录")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run().font.size = Pt(18)
    doc.add_paragraph("第1章 商务文件")
    for index, file_entry in enumerate(selected_files, start=2):
        doc.add_paragraph(f"第{index}章 {file_entry.get('name', '技术应答文件')}")


def _add_commercial(doc: Document) -> None:
    heading = doc.add_paragraph()
    run = heading.add_run("第1章 商务文件")
    run.font.name = "SimHei"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(18)
    doc.add_paragraph("【商务文件内容按招标文件投标文件格式编制，包括投标函、商务偏差表、资质业绩、承诺函等。】")


def _add_product_response(doc: Document, index: int, file_entry: dict) -> None:
    heading = doc.add_paragraph()
    run = heading.add_run(f"第{index + 1}章 {file_entry.get('name', '技术应答文件')}")
    run.font.name = "SimHei"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(18)

    code = file_entry.get("tech_spec_code") or ""
    if code:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(f"技术规范编码：{code}")
        run.bold = True
        run.font.size = Pt(12)

    md_text = ""
    md_path = file_entry.get("md_path", "")
    if md_path and Path(md_path).exists():
        md_text = Path(md_path).read_text(encoding="utf-8", errors="replace")

    _add_section_title(doc, "项目货物技术特性逐条逐项应答文件")
    params_table = _extract_table(md_text, "技术参数特性表")
    _add_table(doc, params_table["header"] or TECH_PARAMS_HEADER, params_table["rows"] or [["", "", "", "", ""] for _ in range(6)])

    _add_section_title(doc, "项目货物组件材料配置表")
    component_table = _extract_table(md_text, "组件材料配置表")
    _add_table(doc, component_table["header"] or COMPONENT_HEADER, component_table["rows"] or [["", "", "", "", "", ""] for _ in range(6)])

    _add_section_title(doc, "技术偏差表")
    _add_table(doc, DEVIATION_HEADER, [["", "", "", "", ""] for _ in range(6)])
    doc.add_paragraph("编制说明：投标人必须进行应答，如有偏差请在本表中说明，无偏差的请在“偏差说明”填写“无偏差”。")


def _add_section_title(doc: Document, title: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(title)
    run.font.name = "SimHei"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(14)


def _add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    column_count = max(len(header), *(len(row) for row in rows), 1)
    table = doc.add_table(rows=1 + len(rows), cols=column_count)
    table.style = "Table Grid"
    for col, value in enumerate(header[:column_count]):
        cell = table.cell(0, col)
        cell.text = ""
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.size = Pt(10)
    for row_index, row in enumerate(rows, start=1):
        for col in range(column_count):
            value = row[col] if col < len(row) else ""
            cell = table.cell(row_index, col)
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            run.font.size = Pt(10)


def _extract_table(md_text: str, keyword: str) -> dict:
    lines = md_text.splitlines()
    for index, line in enumerate(lines):
        if keyword not in line:
            continue
        cursor = index + 1
        while cursor < len(lines) and not lines[cursor].strip().startswith("|"):
            cursor += 1
        if cursor >= len(lines):
            break
        header = _split_row(lines[cursor].strip())
        cursor += 1
        if cursor < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[cursor].strip()):
            cursor += 1
        rows: list[list[str]] = []
        while cursor < len(lines) and lines[cursor].strip().startswith("|"):
            rows.append(_split_row(lines[cursor].strip()))
            cursor += 1
        return {"header": header, "rows": rows}
    return {"header": [], "rows": []}


def _split_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _render_markdown(project: dict, package: dict, selected_files: list[dict]) -> str:
    parts = [f"# 投标文件（{package.get('包名称') or package.get('package_no')}）", ""]
    for label, value in (("项目名称", project.get("项目名称")), ("项目编号", project.get("项目编号")), ("分标名称", project.get("分标名称")), ("分标编号", project.get("分标编号"))):
        if value:
            parts.append(f"{label}：{value}")
    parts.append("")
    for file_entry in selected_files:
        parts.append(f"## {file_entry.get('name')}")
        parts.append(f"技术规范编码：{file_entry.get('tech_spec_code', '')}")
        parts.append("### 项目货物技术特性逐条逐项应答文件")
        parts.append("### 项目货物组件材料配置表")
        parts.append("### 技术偏差表")
        parts.append("")
    return "\n".join(parts)
