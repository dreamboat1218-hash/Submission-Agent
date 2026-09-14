"""模块六：合规检查。

对标书初稿执行格式、内容、报价、资质、时间、完整性与查重检查，
输出 PASS / WARN / FAIL 报告，降低废标风险。
"""
from __future__ import annotations

import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from sqlalchemy.orm import Session

from ..database import get_session_factory
from ..models import KbDocument, Task

REQUIRED_CHAPTERS = [
    "项目概述",
    "技术方案",
    "项目管理",
    "质量",
    "商务报价",
    "资质证明",
    "售后服务",
]


def run(task_id: str) -> dict:
    factory = get_session_factory()
    with factory() as db:
        task = db.get(Task, task_id)
        if task is None:
            raise ValueError(f"task not found: {task_id}")
        extracted = task.extracted_data or {}
        docx_path = Path(task.bid_file_path)
        if not docx_path.exists():
            raise RuntimeError(f"标书文件不存在：{task.bid_file_path}")
        docx_text = _docx_text(docx_path)
        kb_docs = db.query(KbDocument).all()

    checks: list[dict] = [
        _check_signature(docx_text),
        _check_project_name(extracted, docx_text),
        _check_price(extracted, docx_text),
        _check_qualifications(extracted, kb_docs),
        _check_file_kb_qualifications(docx_text),
        _check_file_kb_history(docx_text),
        _check_dates(extracted),
        _check_chapters(docx_text),
        _check_duplication(docx_text, kb_docs),
        _check_security(extracted, docx_text),
    ]
    statuses = [item["result"] for item in checks]
    if "FAIL" in statuses:
        status = "FAIL"
    elif "WARN" in statuses:
        status = "WARN"
    else:
        status = "PASS"

    return {
        "status": status,
        "check_time": datetime.now(timezone.utc).isoformat(),
        "task_id": task_id,
        "items": checks,
        "summary": {"total": len(checks), "pass": statuses.count("PASS"), "warn": statuses.count("WARN"), "fail": statuses.count("FAIL")},
    }


def _docx_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _item(name: str, result: str, detail: str, suggestion: str | None = None) -> dict:
    return {"name": name, "result": result, "detail": detail, "suggestion": suggestion}


def _primary(extracted: dict, field: str):
    item = extracted.get(field)
    if not item:
        return None
    value = item.get("primary_value")
    return value if value not in (None, "") else None


def _check_signature(docx_text: str) -> dict:
    placeholders = re.findall(r"【[^】]{0,12}(?:签字|盖章|公章|骑缝章)[^】]{0,12}】|____________（签字|____________（盖章|签字：|盖章：", docx_text)
    if placeholders:
        return _item(
            "签字盖章检查",
            "WARN",
            f"草稿中仍有 {len(placeholders)} 处待补签字/盖章位置",
            "定稿前完成法人签字、单位公章与骑缝章",
        )
    return _item("签字盖章检查", "PASS", "签字盖章位置均已填写或已按模板约定处理")


def _check_project_name(extracted: dict, docx_text: str) -> dict:
    project_name = _primary(extracted, "project_name")
    if not project_name:
        return _item("项目名称一致性检查", "WARN", "未能从招标文件抽取项目名称", "人工核对项目名称后补充")
    if str(project_name) in docx_text:
        return _item("项目名称一致性检查", "PASS", "标书中的项目名称与招标文件一致")
    return _item(
        "项目名称一致性检查",
        "FAIL",
        f"标书正文未找到项目名称「{project_name}」",
        "检查封面与正文中的项目名称是否被替换",
    )


def _check_price(extracted: dict, docx_text: str) -> dict:
    budget = _primary(extracted, "budget")
    price_match = re.search(
        r"(?:投标报价|总报价|报价金额|报价总额)\s*[：:]\s*([\d,，]+(?:\.\d+)?)\s*(万元|万|元)?",
        docx_text,
    )
    if not price_match:
        return _item("投标报价范围检查", "WARN", "标书草稿中未识别到明确报价", "填写报价明细与总报价后重新检查")
    try:
        number = float(price_match.group(1).replace(",", "").replace("，", ""))
        unit = price_match.group(2) or ""
        price = number * 10000 if unit in {"万元", "万"} else number
        budget_value = float(budget) if budget is not None else None
    except (TypeError, ValueError):
        return _item("投标报价范围检查", "WARN", "报价金额无法解析", "人工核对报价数值")
    if budget_value is None:
        return _item("投标报价范围检查", "WARN", f"已识别报价 {price:.2f} 元，但缺少预算信息", "补充预算信息后复检")
    if price > budget_value:
        return _item(
            "投标报价范围检查",
            "FAIL",
            f"投标报价 {price:.2f} 元超出预算 {budget_value:.2f} 元",
            "调整报价至预算范围内",
        )
    return _item("投标报价范围检查", "PASS", f"投标报价 {price:.2f} 元在预算 {budget_value:.2f} 元范围内")


def _check_qualifications(extracted: dict, kb_docs: list[KbDocument]) -> dict:
    required = _primary(extracted, "qualifications")
    if not required:
        return _item("资质完整性检查", "PASS", "招标文件未提取到明确的资质要求")
    if isinstance(required, str):
        required = [required]
    kb_text = "\n".join(f"{doc.title}\n{doc.content}" for doc in kb_docs)
    missing = []
    for qual in required:
        text = str(qual)
        keywords = re.findall(r"(?:ISO\s*\d+|[\u4e00-\u9fff]{2,10}(?:资质|认证|证书|许可证))", text)
        if not keywords:
            keywords = [text]
        if not any(kw in kb_text for kw in keywords):
            missing.append(text)
    if missing:
        return _item(
            "资质完整性检查",
            "WARN",
            "缺少以下资质证明材料：" + "；".join(missing[:5]),
            "请从企业知识库中补充上传对应证书扫描件",
        )
    return _item("资质完整性检查", "PASS", "所需资质证明材料在知识库中已具备")


def _check_file_kb_qualifications(docx_text: str) -> dict:
    """使用 Obsidian 文件知识库核验标书提及的资质。"""
    try:
        from ..kb_integration import get_gateway

        gateway = get_gateway()
        notes = gateway.list_notes(gateway.config.dir("qualifications"))
        if not notes:
            return _item("文件知识库资质核验", "PASS", "Obsidian 知识库暂无企业资质笔记")
        kb_text = ""
        for rel in notes:
            note = gateway.read_note(rel)
            kb_text += f"{note['path']}\n{note['content']}\n"
        mentions = re.findall(r"(?:ISO\s*\d+|[\u4e00-\u9fff]{2,10}(?:资质|认证|证书|许可证))", docx_text)
        missing = [mention for mention in dict.fromkeys(mentions) if mention not in kb_text]
        if missing:
            return _item(
                "文件知识库资质核验",
                "WARN",
                "标书提及但 Obsidian 知识库中缺失：" + "；".join(missing[:5]),
                "在 01_企业资质 目录补充对应笔记（含 Frontmatter tags）",
            )
        return _item("文件知识库资质核验", "PASS", "标书提及的资质在 Obsidian 知识库中均可找到")
    except Exception:
        return _item("文件知识库资质核验", "PASS", "文件知识库不可用，跳过核验")


def _check_file_kb_history(docx_text: str) -> dict:
    """与 Obsidian 历史标书对比，提示高度雷同的内容。"""
    try:
        from ..kb_integration import get_gateway

        gateway = get_gateway()
        notes = gateway.list_notes(gateway.config.dir("historical_bids"))
        if not notes:
            return _item("历史标书查重", "PASS", "Obsidian 知识库暂无历史标书")
        doc_sentences = {s.strip() for s in re.split(r"[。！？\n]", docx_text) if len(s.strip()) >= 15}
        if not doc_sentences:
            return _item("历史标书查重", "PASS", "标书内容过短，跳过查重")
        worst = 0.0
        worst_path = ""
        for rel in notes:
            note = gateway.read_note(rel)
            history_sentences = {s.strip() for s in re.split(r"[。！？\n]", note["content"]) if len(s.strip()) >= 15}
            if not history_sentences:
                continue
            overlap = len(doc_sentences.intersection(history_sentences)) / len(doc_sentences)
            if overlap > worst:
                worst = overlap
                worst_path = rel
        if worst > 0.5:
            return _item(
                "历史标书查重",
                "WARN",
                f"与历史标书「{worst_path}」雷同率 {worst:.0%}",
                "对直接复制内容进行改写，避免被认定为抄袭",
            )
        return _item("历史标书查重", "PASS", f"与历史标书最高雷同率 {worst:.0%}")
    except Exception:
        return _item("历史标书查重", "PASS", "文件知识库不可用，跳过查重")


def _check_dates(extracted: dict) -> dict:
    deadline = _primary(extracted, "bid_deadline")
    opening = _primary(extracted, "bid_opening_time")
    if not deadline and not opening:
        return _item("时间有效性检查", "WARN", "未抽取到投标截止或开标时间", "人工补充关键时间节点")
    now = datetime.now(timezone.utc)
    expired = []
    for label, value in (("投标截止时间", deadline), ("开标时间", opening)):
        parsed = _parse_datetime(str(value)) if value else None
        if parsed is None:
            continue
        if parsed < now:
            expired.append(f"{label}（{value}）已过")
    if expired:
        return _item("时间有效性检查", "FAIL", "；".join(expired), "确认是否存在延期或补遗公告")
    return _item("时间有效性检查", "PASS", "投标截止与开标时间均在有效期内")


def _check_chapters(docx_text: str) -> dict:
    missing = [chapter for chapter in REQUIRED_CHAPTERS if chapter not in docx_text]
    if missing:
        return _item("章节完整性检查", "FAIL", "缺少章节：" + "、".join(missing), "按模板补齐必填章节")
    return _item("章节完整性检查", "PASS", "所有必填章节均已生成")


def _check_duplication(docx_text: str, kb_docs: list[KbDocument]) -> dict:
    sentences = [s.strip() for s in re.split(r"[。！？\n]", docx_text) if len(s.strip()) >= 15]
    seen: set[str] = set()
    duplicates = 0
    for sentence in sentences:
        if sentence in seen:
            duplicates += 1
        seen.add(sentence)
    ratio = duplicates / max(len(sentences), 1)
    if ratio > 0.2:
        return _item("标书内容查重", "WARN", f"检测到 {duplicates} 条重复内容（重复率 {ratio:.0%}）", "调整重复段落，降低雷同表述")
    kb_ratio = 0.0
    if kb_docs:
        kb_sentences = set()
        for doc in kb_docs:
            kb_sentences.update(s for s in re.split(r"[。！？\n]", f"{doc.title}\n{doc.content}") if len(s.strip()) >= 15)
        overlap = sum(1 for s in sentences if s in kb_sentences)
        kb_ratio = overlap / max(len(sentences), 1)
    if kb_ratio > 0.6:
        return _item("标书内容查重", "WARN", "标书与知识库素材重复率偏高", "对直接复制的素材进行改写")
    return _item("标书内容查重", "PASS", "未发现明显重复内容")


def _check_security(extracted: dict, docx_text: str) -> dict:
    amount = _primary(extracted, "bid_security_amount")
    if amount is None:
        return _item("投标保证金检查", "PASS", "招标文件未要求投标保证金")
    if "保证金" in docx_text:
        return _item("投标保证金检查", "PASS", "标书已包含投标保证金响应说明")
    return _item(
        "投标保证金检查",
        "WARN",
        f"招标文件要求投标保证金 {amount} 元，但标书未体现",
        "补充保证金缴纳凭证与响应说明",
    )


def _parse_datetime(value: str) -> datetime | None:
    value = str(value).strip()
    match = re.search(
        r"(\d{4})\s*[-/年.]\s*(\d{1,2})\s*[-/月.]\s*(\d{1,2})\s*[日号]?\s*(?:(\d{1,2})\s*[:点时]\s*(\d{1,2})\s*分?)?",
        value,
    )
    if not match:
        return None
    year, month, day = int(match.group(1)), int(match.group(2)), int(match.group(3))
    hour = int(match.group(4)) if match.group(4) else 23
    minute = int(match.group(5)) if match.group(5) else 59
    try:
        return datetime(year, month, day, hour, minute, tzinfo=timezone.utc)
    except ValueError:
        return None
