"""国网投标专项包模板的知识库召回与插入。"""
from __future__ import annotations

import logging

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

from ..kb_integration.chapter_mapper import format_knowledge_items

logger = logging.getLogger(__name__)


BID_KNOWLEDGE_HEADINGS = [
    "资质业绩凭证单说明文件",
    "组件材料选用说明",
    "投标人类别",
    "销售业绩（对应招标公告资格条件要求）",
    "检测（试验）报告（对应招标公告资格条件要求）",
    "认证证书（对应招标公告资格条件要求）",
    "许可证（对应招标公告资格条件要求）",
    "工装设备（生产设备）（对应招标公告资格条件要求）",
    "工装设备（试验设备）",
    "其它资质要求（对应招标公告资格条件要求）",
    "人力资源",
    "厂房环境",
    "售后服务及技术支持",
    "组件材料及自主可控水平",
    "创新能力及发明专利",
    "工艺方法",
    "供货能力及产能占用比例",
    "产品制造质量控制",
    "绿色低碳生产及绿色回收",
    "（1）企业名称变更说明及证明材料",
    "（2）其它技术内容",
]

BID_KNOWLEDGE_HEADING_ALIASES = {
    "资质业绩凭证单说明文件": ["资质业绩凭证单"],
    "组件材料选用说明": ["组件材料选用说明", "货物组件材料配置表", "组件材料配置表"],
    "投标人类别": ["投标人类别"],
    "销售业绩（对应招标公告资格条件要求）": ["销售业绩"],
    "检测（试验）报告（对应招标公告资格条件要求）": ["检测检验报告", "检测（试验）报告", "检测报告", "试验报告"],
    "认证证书（对应招标公告资格条件要求）": ["认证证书"],
    "许可证（对应招标公告资格条件要求）": ["许可证"],
    "工装设备（生产设备）（对应招标公告资格条件要求）": ["工装设备（生产设备）", "生产设备"],
    "工装设备（试验设备）": ["工装设备（试验设备）", "试验检测设备", "试验设备"],
    "其它资质要求（对应招标公告资格条件要求）": ["其它资质要求", "其他资质要求"],
    "人力资源": ["人力资源"],
    "厂房环境": ["厂房面积", "厂房环境", "厂房"],
    "售后服务及技术支持": ["售后服务保障及承诺", "售后服务及技术支持", "售后服务"],
    "组件材料及自主可控水平": ["组件材料及自主可控水平", "组件材料"],
    "创新能力及发明专利": ["创新能力及发明专利", "专利"],
    "工艺方法": ["工艺方法"],
    "供货能力及产能占用比例": ["供货能力及产能占用比例", "供货能力"],
    "产品制造质量控制": ["产品制造质量控制"],
    "绿色低碳生产及绿色回收": ["绿色低碳生产及绿色回收", "绿色生产"],
    "（1）企业名称变更说明及证明材料": ["企业名称变更"],
    "（2）其它技术内容": ["其它技术内容", "其他技术内容"],
}


def build_package_knowledge(
    project: dict,
    package: dict,
    industry: str = "",
) -> dict[str, list[dict]]:
    """按国网投标专用规则为包模板固定章节召回知识库内容。"""
    from ..kb_integration import get_gateway
    from ..kb_integration.bid_rules_store import get_bid_chapter_rules
    from ..kb_integration.chapter_mapper import ChapterKnowledgeMapper

    gateway = get_gateway()
    mapper = ChapterKnowledgeMapper(get_bid_chapter_rules())
    tech_names = [
        str(file_entry.get("name", ""))
        for file_entry in package.get("files", [])
        if file_entry.get("kind") == "tech_spec"
    ][:3]
    extracted = {
        "project_name": {"primary_value": project.get("项目名称", "")},
        "technical_requirements": {"primary_value": tech_names},
        "qualifications": {"primary_value": str(package.get("专用资格要求") or "")},
        "performance_requirements": {"primary_value": str(package.get("业绩要求") or "")},
    }
    chapters = [(title, []) for title in BID_KNOWLEDGE_HEADINGS]
    return mapper.build_chapter_map(chapters, extracted, industry or "", gateway)


def inject_knowledge_into_doc(doc: Document, knowledge: dict[str, list[dict]]) -> int:
    """在包模板对应标题下插入“知识库参考内容”，返回实际插入的章节数。"""
    inserted = 0
    for title, items in knowledge.items():
        if not items:
            continue
        aliases = BID_KNOWLEDGE_HEADING_ALIASES.get(title, [title])
        element = _find_heading_element(doc, aliases)
        if element is None:
            continue
        label = _make_paragraph(doc, "知识库参考内容", bold=True)
        element.addnext(label)
        anchor = label
        for line in format_knowledge_items(items).splitlines():
            if not line.strip():
                continue
            body = _make_paragraph(doc, line, bold=False)
            anchor.addnext(body)
            anchor = body
        inserted += 1
    return inserted


def _find_heading_element(doc: Document, aliases: list[str]):
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text or len(text) > 80:
            continue
        if any(alias in text for alias in aliases):
            return paragraph._p
    return None


def _make_paragraph(doc: Document, text: str, bold: bool) -> object:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "宋体"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)
    return paragraph._p
