"""模块三：标书生成。

优先调用 BiaoShu-SKILL 的 ``bid_writer_pipeline.py``；CLI 不可用或执行失败时，
自动降级为内置 DOCX 生成器，保证流水线离线可用。
"""
from __future__ import annotations

import copy
import json
import logging
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph
from sqlalchemy.orm import Session

from ..config import get_settings
from ..database import get_session_factory
from ..models import Task
from . import fallback_extractor
from .industry import biaoshu_code, detect_industry
from .knowledge import KnowledgeService

logger = logging.getLogger(__name__)


class BiaoShuFailed(RuntimeError):
    pass


def run(task_id: str, industry_override: str = "auto", use_kb: bool = True) -> dict:
    settings = get_settings()
    factory = get_session_factory()
    with factory() as db:
        task = db.get(Task, task_id)
        if task is None:
            raise ValueError(f"task not found: {task_id}")
        data = task.extracted_data or {}
        if not data:
            raise RuntimeError("缺少抽取结果，无法生成标书")

    industry = detect_industry(_data_text(data)) if industry_override in {"", "auto"} else industry_override
    kb_context = ""
    kb_docs: list = []
    if use_kb:
        with factory() as db:
            service = KnowledgeService(db)
            query = " ".join(
                str(v.get("primary_value", ""))
                for v in data.values()
                if isinstance(v, dict) and v.get("primary_value")
            )[:200]
            kb_context = service.retrieve_context(query or "企业资质 项目业绩", top_k=5)
            _, kb_docs = service.list_documents(limit=200)
        file_kb_context = _file_knowledge_context(task_id, data, industry)
        if file_kb_context:
            kb_context = (kb_context + "\n\n" + file_kb_context).strip() if kb_context else file_kb_context

    output_dir = Path(task.work_dir) / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    target_docx = output_dir / f"bid_{task_id[:8]}.docx"

    provider = "builtin"
    try:
        if settings.biaoshu_skill_path:
            source = _run_biaoshu(task, data, industry, output_dir, target_docx)
            provider = "biaoshu"
            logger.info("BiaoShu-SKILL 生成成功 task=%s", task_id)
    except Exception as exc:
        logger.warning("BiaoShu-SKILL 执行失败，降级为内置生成器：%s", exc)

    if provider == "builtin":
        _run_builtin(task, data, industry, kb_context, output_dir, target_docx, kb_docs)

    intermediates = output_dir / "intermediates"
    intermediates.mkdir(exist_ok=True)
    with factory() as db:
        task = db.get(Task, task_id)
        combined = Path(task.work_dir) / "combined.md"
        if combined.exists():
            shutil.copy2(combined, intermediates / "tender_combined.md")
        (intermediates / "extracted_data.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        (intermediates / "kb_context.md").write_text(kb_context, encoding="utf-8")

    return {"bid_file_path": str(target_docx), "industry": industry, "provider": provider}


def _data_text(data: dict) -> str:
    parts = []
    for field in fallback_extractor.FIELD_LABELS:
        item = data.get(field)
        if item:
            value = item.get("primary_value")
            if isinstance(value, list):
                parts.extend(str(v) for v in value)
            elif value is not None:
                parts.append(str(value))
    return "\n".join(parts)


def _file_knowledge_context(task_id: str, data: dict, industry: str) -> str:
    """从 Obsidian 知识库组装标书生成素材，注入相似案例、资质、模板与政策。"""
    try:
        from ..kb_integration import get_gateway

        gateway = get_gateway()
        context: dict | None = None
        with get_session_factory()() as db:
            task = db.get(Task, task_id)
            if task is not None:
                json_path = Path(task.work_dir) / "output" / "intermediates" / "knowledge_context.json"
                if json_path.exists():
                    try:
                        context = json.loads(json_path.read_text(encoding="utf-8"))
                    except json.JSONDecodeError:
                        context = None
        if not context:
            description = _value(data, "project_name", "") or _data_text(data)[:500]
            context = gateway.get_context_for_bid(description, industry or "")
        parts: list[str] = []
        for group, label in (
            ("similar_projects", "相似历史项目"),
            ("qualifications", "企业资质"),
            ("templates", "投标模板"),
            ("policies", "行业政策"),
        ):
            items = context.get(group) or []
            if not items:
                continue
            block = [f"## {label}"]
            for item in items[:3]:
                block.append(f"### {item.get('path')}")
                block.append(str(item.get("content", ""))[:1500])
            parts.append("\n".join(block))
        return "\n\n".join(parts)
    except Exception as exc:
        logger.warning("Obsidian 知识库上下文注入失败：%s", exc)
        return ""


def _run_biaoshu(task: Task, data: dict, industry: str, output_dir: Path, target_docx: Path) -> Path:
    settings = get_settings()
    skill_root = Path(settings.biaoshu_skill_path)
    pipeline_script = skill_root / "scripts" / "bid_writer_pipeline.py"
    if not pipeline_script.exists():
        raise BiaoShuFailed(f"找不到 {pipeline_script}")

    run_dir = Path(task.work_dir) / "biaoshu_run"
    run_dir.mkdir(parents=True, exist_ok=True)
    combined = Path(task.work_dir) / "combined.md"
    # BiaoShu 脚本以自己的目录为 cwd 执行，路径必须绝对，否则解析不到输入文件
    input_txt = (run_dir / "tender_input.txt").resolve()
    if combined.exists():
        shutil.copy2(combined, input_txt)
    else:
        input_txt.write_text(_data_text(data), encoding="utf-8")

    out_dir = (run_dir / "out").resolve()
    command = [
        sys.executable,
        str(pipeline_script),
        str(input_txt),
        "--output-dir",
        str(out_dir),
        "--industry",
        biaoshu_code(industry),
        "--format-standard",
        "government",
    ]
    try:
        result = subprocess.run(
            command,
            cwd=str(skill_root),
            capture_output=True,
            text=True,
            timeout=settings.biaoshu_timeout,
        )
    except subprocess.TimeoutExpired as exc:
        raise BiaoShuFailed(f"BiaoShu-SKILL 执行超时（{settings.biaoshu_timeout}s）") from exc
    if result.returncode != 0:
        raise BiaoShuFailed(f"BiaoShu-SKILL 退出码 {result.returncode}：{result.stderr[-1000:]}")

    candidates = sorted(run_dir.rglob("*.docx"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not candidates:
        raise BiaoShuFailed("BiaoShu-SKILL 未输出 DOCX 文件")
    shutil.copy2(candidates[0], target_docx)
    return target_docx


def _run_builtin(
    task: Task,
    data: dict,
    industry: str,
    kb_context: str,
    output_dir: Path,
    target_docx: Path,
    kb_docs: list | None = None,
) -> Path:
    primary = _bid_format_primary(data)
    source_path = str(primary.get("source_path", "") or "")
    if source_path and Path(source_path).suffix.lower() == ".docx" and Path(source_path).exists():
        try:
            return _build_raw_from_source_doc(data, industry, Path(source_path), output_dir, target_docx, kb_docs)
        except Exception as exc:
            logger.warning("源 DOCX 格式模板复制失败，降级为 Markdown 重建：%s", exc)

    doc = Document()
    _setup_document(doc)

    project_name = _value(data, "project_name", "未命名项目")
    bid_number = _value(data, "bid_number", "")
    _add_cover(doc, project_name, bid_number, industry)
    doc.add_page_break()

    if _build_raw_template_doc(doc, data, industry, kb_docs):
        markdown = _raw_template_markdown(data)
        (output_dir / "bid_draft.md").write_text(markdown, encoding="utf-8")
        doc.save(target_docx)
        return target_docx

    chapters = _build_chapter_plan(data, industry, kb_context)
    chapter_knowledge = _build_chapter_knowledge(chapters, data, industry)
    chapters = _attach_chapter_knowledge(chapters, chapter_knowledge)
    _add_chapter_list(doc, chapters)
    doc.add_page_break()

    for index, (title, sections) in enumerate(chapters, start=1):
        _add_chapter(doc, index, title, sections)
        if index < len(chapters):
            doc.add_page_break()
    _append_kb_attachments(doc, kb_docs, "商务文件")
    _append_kb_attachments(doc, kb_docs, "技术文件")

    if chapter_knowledge:
        report_dir = output_dir / "intermediates"
        report_dir.mkdir(parents=True, exist_ok=True)
        report_lines = ["# 知识库章节引用报告", ""]
        for title, items in chapter_knowledge.items():
            report_lines.append(f"## {title}")
            for item in items:
                report_lines.append(f"- {item.get('path')}（{item.get('source', '')}，相关度 {item.get('score', '')}）")
            report_lines.append("")
        (report_dir / "knowledge_chapter_report.md").write_text("\n".join(report_lines), encoding="utf-8")

    markdown = _render_markdown(chapters, project_name, industry)
    (output_dir / "bid_draft.md").write_text(markdown, encoding="utf-8")
    doc.save(target_docx)
    return target_docx


def _find_body_element(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if text in paragraph.text:
            return paragraph._p
    return None


def _build_raw_from_source_doc(
    data: dict,
    industry: str,
    source_path: Path,
    output_dir: Path,
    target_docx: Path,
    kb_docs: list | None = None,
) -> Path:
    """直接复制招标文件原 DOCX 的商务/技术部分，完整保留表格、合并单元格与页面版式。"""
    shutil.copy2(source_path, target_docx)
    doc = Document(target_docx)

    commercial_el = _find_body_element(doc, "商务部分（单独成册）")
    technical_el = _find_body_element(doc, "技术部分（单独成册）")
    if commercial_el is None or technical_el is None:
        raise RuntimeError("未在源 DOCX 中找到商务部分/技术部分标记")

    end_el = _find_body_element(doc, "投标保证保险（责任要求）")
    if end_el is None:
        end_el = _find_body_element(doc, "投标保证金银行保函（格式）")

    cover_doc = Document()
    _setup_document(cover_doc)
    project_name = _value(data, "project_name", "未命名项目")
    bid_number = _value(data, "bid_number", "")
    _add_cover(cover_doc, project_name, bid_number, industry)
    cover_doc.add_page_break()
    _add_chapter_list(cover_doc, [("商务文件", []), ("技术文件", [])])
    cover_doc.add_page_break()

    body = doc.element.body
    cover_elements = [
        copy.deepcopy(child)
        for child in list(cover_doc.element.body)
        if child.tag != qn("w:sectPr")
    ]
    first_new = None
    for element in cover_elements:
        commercial_el.addprevious(element)
        if first_new is None:
            first_new = element

    for child in list(body):
        if child is first_new:
            break
        body.remove(child)

    if end_el is not None:
        removing = False
        for child in list(body):
            if child is end_el:
                removing = True
            if removing and child.tag != qn("w:sectPr"):
                body.remove(child)

    page_break_paragraph = doc.add_paragraph()
    page_break_paragraph.add_run().add_break(WD_BREAK.PAGE)
    technical_el.addprevious(page_break_paragraph._p)

    _append_kb_attachments(doc, kb_docs, "商务文件", anchor=page_break_paragraph._p)
    _append_kb_attachments(doc, kb_docs, "技术文件")

    (output_dir / "bid_draft.md").write_text(_raw_template_markdown(data), encoding="utf-8")
    doc.save(target_docx)
    return target_docx


def _append_kb_attachments(doc: Document, kb_docs: list | None, chapter: str, anchor=None) -> None:
    """把知识库附件内容插入标书对应章节；PDF 渲染页与图片附件以图片形式嵌入。"""
    docs = [
        kb
        for kb in (kb_docs or [])
        if (getattr(kb, "metadata_json", None) or {}).get("chapter") == chapter
    ]
    if not docs:
        return

    existing = set(doc.element.body)
    for kb in docs:
        heading = doc.add_paragraph()
        run = heading.add_run(f"【{chapter}附件】{kb.title}")
        run.bold = True
        run.font.name = "SimHei"
        run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.size = Pt(14)

        content = str(kb.content or "").strip()
        if content:
            for line in content.splitlines()[:200]:
                line = line.strip()
                if line:
                    paragraph = doc.add_paragraph(line)
                    paragraph.paragraph_format.first_line_indent = Pt(24)

        images = (kb.metadata_json or {}).get("images", []) or []
        for image_path in images[:12]:
            if Path(str(image_path)).exists():
                try:
                    doc.add_picture(str(image_path), width=Inches(5.8))
                except Exception:
                    logger.warning("知识库图片嵌入失败：%s", image_path)

    if anchor is not None:
        new_children = [el for el in list(doc.element.body) if el not in existing and el.tag != qn("w:sectPr")]
        for element in new_children:
            anchor.addprevious(element)


def _bid_format_primary(data: dict) -> dict:
    fmt = data.get("bid_format", {})
    primary = fmt.get("primary_value") if isinstance(fmt, dict) else None
    return primary if isinstance(primary, dict) else {}


def _build_raw_template_doc(doc: Document, data: dict, industry: str, kb_docs: list | None = None) -> bool:
    """按招标文件投标文件格式原文生成：只包含商务文件与技术文件两册。"""
    primary = _bid_format_primary(data)
    commercial_raw = str(primary.get("commercial_raw", "") or "").strip()
    technical_raw = str(primary.get("technical_raw", "") or "").strip()
    if not commercial_raw and not technical_raw:
        return False

    chapters: list[tuple[str, str]] = []
    if commercial_raw:
        chapters.append(("商务文件", commercial_raw))
    if technical_raw:
        chapters.append(("技术文件", technical_raw))

    _add_chapter_list(doc, [(title, []) for title, _ in chapters])
    doc.add_page_break()
    for title, raw in chapters:
        _add_raw_section(doc, title, raw)
        doc.add_page_break()
    _append_kb_attachments(doc, kb_docs, "商务文件")
    _append_kb_attachments(doc, kb_docs, "技术文件")
    return True


def _raw_template_markdown(data: dict) -> str:
    primary = _bid_format_primary(data)
    parts: list[str] = []
    for title, key in (("商务文件", "commercial_raw"), ("技术文件", "technical_raw")):
        raw = str(primary.get(key, "") or "").strip()
        if raw:
            parts.append(f"# {title}\n\n{raw}")
    return "\n\n".join(parts)


def _add_raw_section(doc: Document, title: str, raw: str) -> None:
    heading = doc.add_paragraph()
    run = heading.add_run(title)
    run.font.name = "SimHei"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(20)
    _add_markdown_block(doc, raw)


def _add_markdown_block(doc: Document, text: str) -> None:
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[index + 1]):
            header = _split_table_row(stripped)
            rows: list[list[str]] = []
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_split_table_row(lines[index].strip()))
                index += 1
            _add_markdown_table(doc, header, rows)
            continue
        if stripped.startswith("#"):
            level = min(len(stripped) - len(stripped.lstrip("#")), 3)
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(stripped.lstrip("# ").strip())
            run.font.name = "SimHei"
            run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
            run.font.size = Pt(18 if level == 1 else 14 if level == 2 else 12)
            index += 1
            continue
        if re.match(r"^\d+[\.、]", stripped) and len(stripped) <= 60:
            paragraph = doc.add_paragraph()
            run = paragraph.add_run(stripped)
            run.bold = True
            run.font.size = Pt(12)
            index += 1
            continue
        paragraph = doc.add_paragraph(stripped)
        paragraph.paragraph_format.first_line_indent = Pt(24)
        index += 1


def _split_table_row(line: str) -> list[str]:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return cells


def _add_markdown_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    column_count = max(len(header), *(len(row) for row in rows), 1)
    table = doc.add_table(rows=1 + len(rows), cols=column_count)
    table.style = "Table Grid"
    for col_index, value in enumerate(header[:column_count]):
        _set_cell(table.cell(0, col_index), value, bold=True)
    for row_index, row in enumerate(rows, start=1):
        for col_index in range(column_count):
            value = row[col_index] if col_index < len(row) else ""
            _set_cell(table.cell(row_index, col_index), value)


def _set_cell(cell, value: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.size = Pt(9)


def _value(data: dict, field: str, default: Any = "") -> Any:
    item = data.get(field)
    if not item:
        return default
    value = item.get("primary_value")
    return value if value not in (None, "", []) else default


def _setup_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for section in doc.sections:
        section.top_margin = Pt(72)
        section.bottom_margin = Pt(72)
        section.left_margin = Pt(90)
        section.right_margin = Pt(90)


def _add_cover(doc: Document, project_name: str, bid_number: str, industry: str) -> None:
    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("投 标 文 件")
    run.font.name = "SimHei"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(36)
    doc.add_paragraph()
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"（{project_name}）")
    run.font.size = Pt(16)
    info_lines = [f"招标编号：{bid_number}", f"行业类型：{industry}", "投标人：____________（盖章）", "法定代表人或授权代表：____________（签字）", "日期：____________"]
    for _ in range(4):
        doc.add_paragraph()
    for line in info_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line).font.size = Pt(14)


def _add_chapter_list(doc: Document, chapters: list[tuple[str, list[tuple[str, str]]]]) -> None:
    heading = doc.add_paragraph("目 录")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run().font.size = Pt(18)
    for index, (name, _) in enumerate(chapters, start=1):
        doc.add_paragraph(f"第{index}章 {name}")


def _build_chapter_plan(data: dict, industry: str, kb_context: str) -> list[tuple[str, list[tuple[str, str]]]]:
    """根据招标文件中识别到的投标文件格式，动态生成章节结构。"""
    fmt = data.get("bid_format", {})
    primary = fmt.get("primary_value") if isinstance(fmt, dict) else None
    sections = primary.get("sections", []) if isinstance(primary, dict) else []
    mapped = _map_format_sections(sections)
    if not mapped:
        return [
            ("项目概述", _chapter_project_overview(data)),
            ("公司简介与项目业绩", _chapter_company(data, kb_context)),
            ("技术方案", _chapter_technical(data, industry)),
            ("项目管理与实施", _chapter_management(data)),
            ("质量承诺与售后服务", _chapter_quality(data)),
            ("商务报价", _chapter_price(data)),
            ("资质证明", _chapter_qualifications(data, kb_context)),
            ("响应性说明与格式符合性", _chapter_format(data)),
        ]

    chapters: list[tuple[str, list[tuple[str, str]]]] = []
    seen: set[str] = set()
    for title, kind in mapped:
        if title in seen:
            continue
        seen.add(title)
        if kind == "commercial":
            chapters.append((title, _chapter_commercial(data, kb_context)))
        elif kind == "technical":
            chapters.append((title, _chapter_technical_file(data, industry)))
        elif kind == "price":
            chapters.append((title, _chapter_price_file(data)))
        elif kind == "qualification":
            chapters.append((title, _chapter_qualification_file(data, kb_context)))
        elif kind == "letter":
            chapters.append((title, _chapter_bid_letter(data)))
        elif kind == "table":
            chapters.append((title, _chapter_format_table(title, data)))
        else:
            chapters.append((title, [("编制说明", f"按招标文件「{title}」格式要求编制。\n【内容占位符】")]))
    chapters.append(("响应性说明与格式符合性", _chapter_format(data)))
    return chapters


def _build_chapter_knowledge(chapters: list[tuple[str, list]], data: dict, industry: str) -> dict[str, list[dict]]:
    """按章节从 Obsidian 知识库召回内容（方案 A 规则 + 方案 B 语义检索）。"""
    try:
        from ..kb_integration import get_gateway
        from ..kb_integration.chapter_mapper import ChapterKnowledgeMapper

        gateway = get_gateway()
        mapper = ChapterKnowledgeMapper()
        file_map = mapper.build_chapter_map(chapters, data, industry, gateway)
        db_map = _build_db_chapter_knowledge(chapters, data, industry, mapper)
        merged: dict[str, list[dict]] = {}
        for title, items in file_map.items():
            merged[title] = list(items)
        for title, items in db_map.items():
            existing = merged.setdefault(title, [])
            seen = {item.get("path") for item in existing}
            for item in items:
                if item.get("path") not in seen:
                    existing.append(item)
                    seen.add(item.get("path"))
            merged[title] = sorted(existing, key=lambda item: float(item.get("score", 0) or 0), reverse=True)
        return merged
    except Exception as exc:
        logger.warning("章节知识映射失败（不影响生成）：%s", exc)
        return {}


def _build_db_chapter_knowledge(
    chapters: list[tuple[str, list]],
    data: dict,
    industry: str,
    mapper,
) -> dict[str, list[dict]]:
    """从原有资料库（数据库知识库）按章节补充召回，与 Obsidian 知识库互补。"""
    from ..database import get_session_factory
    from .knowledge import KnowledgeService

    result: dict[str, list[dict]] = {}
    with get_session_factory()() as db:
        service = KnowledgeService(db)
        for title, _ in chapters:
            rule = mapper.match_chapter(title)
            query = mapper.build_query(title, rule or {}, data, industry) if rule else title
            try:
                hits = service.search(query or title, top_k=3)
            except Exception:
                hits = []
            items = [
                {
                    "path": f"资料库/{hit.get('title', '')}",
                    "score": float(hit.get("score", 0) or 0),
                    "content": str(hit.get("content", ""))[:2000],
                    "metadata": {"doc_type": hit.get("doc_type")},
                    "source": "资料库",
                }
                for hit in hits
            ]
            if items:
                result[title] = items
    return result


def _attach_chapter_knowledge(
    chapters: list[tuple[str, list[tuple[str, str]]]],
    chapter_knowledge: dict[str, list[dict]],
) -> list[tuple[str, list[tuple[str, str]]]]:
    from ..kb_integration.chapter_mapper import format_knowledge_items

    result: list[tuple[str, list[tuple[str, str]]]] = []
    for title, sections in chapters:
        items = chapter_knowledge.get(title) or []
        if items:
            sections = list(sections) + [("知识库参考内容", format_knowledge_items(items))]
        result.append((title, sections))
    return result


def _map_format_sections(sections: list) -> list[tuple[str, str]]:
    mapped: list[tuple[str, str]] = []
    for section in sections:
        title = str(section).strip()
        if any(k in title for k in ("商务文件", "商务标", "商务部分")):
            mapped.append(("商务文件", "commercial"))
        elif any(k in title for k in ("技术文件", "技术标", "技术部分")):
            mapped.append(("技术文件", "technical"))
        elif any(k in title for k in ("价格文件", "报价文件", "报价书")):
            mapped.append(("价格文件", "price"))
        elif any(k in title for k in ("资格文件", "资格预审文件")):
            mapped.append(("资格文件", "qualification"))
        elif "投标函" in title:
            mapped.append(("投标函", "letter"))
        elif any(k in title for k in ("开标一览表", "报价明细表", "偏差表", "保证金明细表")):
            mapped.append((title, "table"))
        else:
            mapped.append((title, "other"))
    return mapped


def _add_chapter(doc: Document, index: int, title: str, sections: list[tuple[str, str]]) -> None:
    heading = doc.add_paragraph()
    run = heading.add_run(f"第{index}章 {title}")
    run.font.name = "SimHei"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    run.font.size = Pt(18)
    for section_title, content in sections:
        sub = doc.add_paragraph()
        sub_run = sub.add_run(section_title)
        sub_run.font.name = "SimHei"
        sub_run.element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        sub_run.font.size = Pt(14)
        for line in content.splitlines():
            p = doc.add_paragraph(line)
            p.paragraph_format.first_line_indent = Pt(24)


def _chapter_project_overview(data: dict) -> list[tuple[str, str]]:
    rows = [
        ("项目名称", _value(data, "project_name", "")),
        ("招标编号", _value(data, "bid_number", "")),
        ("招标人", _value(data, "tenderer", "")),
        ("项目地点", _value(data, "project_location", "")),
        ("项目规模", _value(data, "project_scale", "")),
        ("项目预算", _value(data, "budget", "")),
        ("投标截止时间", _value(data, "bid_deadline", "")),
        ("开标时间", _value(data, "bid_opening_time", "")),
        ("工期要求", _value(data, "delivery_period", "")),
    ]
    lines = "\n".join(f"1. {label}：{value if value not in (None, '') else '____'}" for label, value in rows)
    return [("项目背景与理解", lines), ("项目目标", "我司充分理解本项目的建设目标与业务要求，将严格按照招标文件规定的范围、标准与工期完成交付，确保项目质量与用户满意度。")]


def _chapter_company(data: dict, kb_context: str) -> list[tuple[str, str]]:
    intro = kb_context[:1200] if kb_context else "【公司简介占位符：请从企业知识库或公司资料中补充公司规模、主营业务、核心能力。】"
    performance = _value(data, "performance_requirements", "【业绩要求占位符：请补充近三年类似项目业绩。】")
    return [
        ("公司简介", str(intro)),
        ("类似项目业绩", f"针对招标文件的业绩要求：{performance}\n\n【项目业绩表占位符：项目名称、合同金额、完成时间、客户名称、验收情况。】"),
    ]


def _chapter_technical(data: dict, industry: str) -> list[tuple[str, str]]:
    requirement = _value(data, "technical_requirements", "【技术要求占位符：请按招标文件补充。】")
    return [
        ("总体方案", f"本投标方案针对「{_value(data, 'project_name', '本项目')}」编制，结合{industry}行业特点设计总体实施方案。"),
        ("技术需求响应", f"招标文件技术要求：{requirement}\n\n响应说明：我司将逐条响应全部技术要求，并提供详细的设备清单、配置方案与测试报告。"),
        ("关键交付内容", "【技术架构图占位符】\n【功能模块清单占位符】\n【实施计划表占位符】"),
    ]


def _chapter_management(data: dict) -> list[tuple[str, str]]:
    period = _value(data, "delivery_period", "【工期占位符】")
    quality = _value(data, "quality_standard", "【质量标准占位符】")
    return [
        ("实施组织", "成立专项项目部，配置项目经理、技术负责人、实施工程师和质控专员，明确岗位职责与汇报机制。"),
        ("进度计划", f"本项目建设/服务期为：{period}。\n【项目进度计划表占位符】"),
        ("质量与风险控制", f"质量执行标准：{quality}\n风险控制：建立风险登记册，每周评审，重大风险 24 小时内上报并启动预案。"),
    ]


def _chapter_quality(data: dict) -> list[tuple[str, str]]:
    return [
        ("质量承诺", f"我司承诺项目验收达到：{_value(data, 'quality_standard', '国家及行业现行标准')}。"),
        ("售后服务", "提供 7×24 小时服务热线，紧急故障 15 分钟响应、2 小时内到场处理；质保期内免费维护。"),
        ("培训计划", "【培训计划表占位符：培训内容、对象、方式、时间、地点。】"),
    ]


def _chapter_price(data: dict) -> list[tuple[str, str]]:
    budget = _value(data, "budget", "")
    try:
        suggested = round(float(budget) / 10000, 2) if budget else ""
        hint = f"建议报价区间：{suggested}万元（预算 {budget} 元，请结合成本测算后填写）" if suggested else ""
    except (TypeError, ValueError):
        hint = "【请按企业成本测算填写报价】"
    return [
        ("报价说明", f"本报价书严格响应招标文件要求。{hint}"),
        ("报价清单", "【报价明细表占位符：序号、项目、单位、数量、单价、合价。】"),
        ("付款方式响应", f"付款方式：{_value(data, 'payment_terms', '【按招标文件约定填写】')}"),
    ]


def _chapter_qualifications(data: dict, kb_context: str) -> list[tuple[str, str]]:
    quals = _value(data, "qualifications", [])
    if isinstance(quals, list) and quals:
        lines = "\n".join(f"- {q}" for q in quals)
    else:
        lines = str(quals) if quals else "【资质要求占位符】"
    kb_hint = "知识库中的资质资料：" + kb_context[:600] if kb_context else ""
    return [
        ("资质响应", f"招标文件资质要求：\n{lines}\n\n{kb_hint}"),
        ("资质证明材料", "【资质证书扫描件占位符：营业执照、ISO认证、专业承包资质等。】"),
    ]


def _chapter_format(data: dict) -> list[tuple[str, str]]:
    formats = _value(data, "bid_format_requirements", [])
    if isinstance(formats, list) and formats:
        lines = "\n".join(f"- {f}" for f in formats)
    else:
        lines = str(formats) if formats else "【格式要求占位符】"
    fmt = data.get("bid_format", {})
    primary = fmt.get("primary_value") if isinstance(fmt, dict) else None
    extra_lines: list[str] = []
    if isinstance(primary, dict):
        if primary.get("requirements"):
            extra_lines.append("招标文件格式要求：" + "、".join(str(item) for item in primary["requirements"]))
        if primary.get("templates"):
            extra_lines.append("招标文件格式模板：" + "、".join(str(item) for item in primary["templates"]))
    if extra_lines:
        lines = lines + "\n" + "\n".join(extra_lines) if lines else "\n".join(extra_lines)
    return [
        ("格式符合性说明", f"招标文件格式要求：\n{lines}\n\n我司将按招标文件要求完成装订、密封、份数与签字盖章。"),
        ("签字盖章", "【法定代表人签字】\n【单位公章】\n【骑缝章】"),
    ]


def _chapter_commercial(data: dict, kb_context: str) -> list[tuple[str, str]]:
    return [
        ("公司简介与业绩", _chapter_company(data, kb_context)[0][1]),
        ("资质响应", _chapter_qualifications(data, kb_context)[0][1]),
        ("投标保证金响应", f"保证金金额：{_value(data, 'bid_security_amount', '【按招标文件填写】')} 元；缴纳截止：{_value(data, 'bid_security_deadline', '【按招标文件填写】')}。"),
        ("商务偏差表", "【商务偏差表占位符：序号、招标文件条款、投标响应内容、偏差说明。】"),
    ]


def _chapter_technical_file(data: dict, industry: str) -> list[tuple[str, str]]:
    return [
        ("总体方案", _chapter_technical(data, industry)[0][1]),
        ("技术需求响应", _chapter_technical(data, industry)[1][1]),
        ("项目管理与实施", _chapter_management(data)[0][1]),
        ("质量承诺与售后服务", _chapter_quality(data)[0][1]),
        ("技术偏差表", "【技术偏差表占位符：序号、招标技术要求、投标响应参数、偏差说明。】"),
    ]


def _chapter_price_file(data: dict) -> list[tuple[str, str]]:
    return [
        ("报价说明", _chapter_price(data)[0][1]),
        ("开标一览表", "【开标一览表占位符：项目名称、投标总报价、供货期、质量标准。】"),
        ("报价明细表", "【报价明细表占位符：序号、货物名称、规格型号、单位、数量、未含税单价、税率、含税合价。】"),
        ("付款方式响应", f"付款方式：{_value(data, 'payment_terms', '【按招标文件约定填写】')}"),
    ]


def _chapter_qualification_file(data: dict, kb_context: str) -> list[tuple[str, str]]:
    return [
        ("资格响应", _chapter_qualifications(data, kb_context)[0][1]),
        ("证明材料", _chapter_qualifications(data, kb_context)[1][1]),
        ("业绩凭证", "【资质业绩凭证单/业绩合同扫描件占位符】"),
    ]


def _chapter_bid_letter(data: dict) -> list[tuple[str, str]]:
    return [
        (
            "投标函",
            f"致：{_value(data, 'tenderer', '招标人')}\n\n"
            f"我方已仔细阅读《{_value(data, 'project_name', '本项目')}》招标文件（编号：{_value(data, 'bid_number', '____')}），"
            "愿意按招标文件要求承担本项目，投标总报价：【____】元。\n\n"
            "投标人：____________（盖章）\n法定代表人或授权代表：____________（签字）\n日期：____年__月__日",
        ),
        ("承诺事项", "【投标人承诺函占位符：无违法失信、无关联关系、诚信投标等承诺。】"),
    ]


def _chapter_format_table(title: str, data: dict) -> list[tuple[str, str]]:
    hints = {
        "开标一览表": "项目名称、投标总报价、供货期、质量标准。",
        "报价明细表": "序号、货物名称、规格型号、单位、数量、未含税单价、增值税税率、含税合价。",
        "偏差表": "序号、招标文件条款、投标响应内容、偏差说明。",
        "保证金明细表": "保证金序号、分标名称、缴纳方式、金额、凭证编号。",
    }
    hint = next((value for key, value in hints.items() if key in title), "按招标文件要求填写。")
    return [(title, f"【{title}占位符：{hint}】")]


def _render_markdown(chapters: list[tuple[str, list[tuple[str, str]]]], project_name: str, industry: str) -> str:
    parts = [f"# 投标文件：{project_name}", f"\n> 行业：{industry}\n"]
    for index, (title, sections) in enumerate(chapters, start=1):
        parts.append(f"\n## 第{index}章 {title}\n")
        for section_title, content in sections:
            parts.append(f"### {section_title}\n\n{content}")
    return "\n".join(parts)
