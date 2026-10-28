from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path

from docx import Document

from app.database import get_session_factory
from app.models import Task
from app.modules import compliance


def _make_docx(path: Path, price_line: str | None = None) -> None:
    doc = Document()
    doc.add_heading("投标文件", 0)
    doc.add_paragraph("第一章 项目概述")
    doc.add_paragraph("项目名称：智慧交通建设项目")
    doc.add_paragraph("第二章 技术方案")
    doc.add_paragraph("平台支持500路视频接入。")
    doc.add_paragraph("第三章 项目管理")
    doc.add_paragraph("第四章 质量承诺与售后服务")
    doc.add_paragraph("第五章 商务报价")
    if price_line:
        doc.add_paragraph(price_line)
    doc.add_paragraph("第六章 资质证明")
    doc.add_paragraph("第七章 响应性说明与格式符合性")
    doc.add_paragraph("第八章 投标保证金响应说明")
    doc.save(path)


def _task_with(path: Path, budget: float) -> Task:
    deadline = (datetime.now() + timedelta(days=30)).strftime("%Y-%m-%d 17:00")
    opening = (datetime.now() + timedelta(days=38)).strftime("%Y-%m-%d 09:30")
    factory = get_session_factory()
    with factory() as db:
        task = Task(
            user_id="test",
            status="completed",
            original_file_name="tender.md",
            work_dir=str(path.parent),
            bid_file_path=str(path),
            extracted_data={
                "project_name": {"primary_value": "智慧交通建设项目", "confidence": 0.99, "evidence": {}},
                "budget": {"primary_value": budget, "confidence": 0.99, "evidence": {}},
                "bid_deadline": {"primary_value": deadline, "confidence": 0.99, "evidence": {}},
                "bid_opening_time": {"primary_value": opening, "confidence": 0.99, "evidence": {}},
                "bid_security_amount": {"primary_value": 80000.0, "confidence": 0.99, "evidence": {}},
                "qualifications": {
                    "primary_value": ["ISO9001质量管理体系认证", "电子与智能化工程专业承包二级"],
                    "confidence": 0.9,
                    "evidence": {},
                },
            },
        )
        db.add(task)
        db.commit()
        db.refresh(task)
        return task


def test_compliance_rejects_over_budget(tmp_path):
    docx = tmp_path / "bid.docx"
    _make_docx(docx, price_line="投标报价：580万元")
    task = _task_with(docx, 5680000.0)
    report = compliance.run(task.id)
    assert report["status"] == "FAIL"
    price_item = next(item for item in report["items"] if item["name"] == "投标报价范围检查")
    assert price_item["result"] == "FAIL"


def test_compliance_accepts_in_budget(tmp_path):
    docx = tmp_path / "bid.docx"
    _make_docx(docx, price_line="投标报价：560万元")
    task = _task_with(docx, 5680000.0)
    report = compliance.run(task.id)
    assert report["status"] in {"PASS", "WARN"}
    price_item = next(item for item in report["items"] if item["name"] == "投标报价范围检查")
    assert price_item["result"] == "PASS"
