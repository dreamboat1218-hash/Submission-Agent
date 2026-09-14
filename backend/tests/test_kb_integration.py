"""Obsidian 知识库集成模块单元测试。"""
from __future__ import annotations

from pathlib import Path

import pytest

from app.kb_integration.config import KnowledgeBaseConfig
from app.kb_integration.context_builder import ContextBuilder, extract_keywords
from app.kb_integration.file_handler import (
    extract_tags,
    extract_wikilinks,
    list_notes,
    parse_note,
    read_note,
    render_note,
    safe_relative_path,
    write_note,
)
from app.kb_integration.gateway import UnifiedKnowledgeGateway, reset_gateway
from app.kb_integration.chapter_mapper import ChapterKnowledgeMapper, format_knowledge_items
from app.kb_integration.pinakes_adapter import PinakesAdapter
from app.kb_integration.tag_handler import search_by_tag


def _config(tmp_path: Path) -> KnowledgeBaseConfig:
    return KnowledgeBaseConfig(
        root_path=tmp_path / "kb",
        pinakes_bin=tmp_path / "missing-pinakes",
        max_search_results=5,
        search_timeout=2.0,
    )


def test_parse_and_render_frontmatter() -> None:
    text = "---\ntitle: 测试笔记\ntags:\n  - 企业资质\n  - ISO9001\nindustry: IT\n---\n\n# 正文\n内容"
    metadata, content = parse_note(text)
    assert metadata["title"] == "测试笔记"
    assert metadata["tags"] == ["企业资质", "ISO9001"]
    assert "正文" in content

    rendered = render_note("正文内容", {"title": "新笔记", "tags": ["资质"]})
    assert rendered.startswith("---\n")
    assert "title: 新笔记" in rendered
    assert rendered.endswith("正文内容")

    parsed_meta, parsed_content = parse_note(rendered)
    assert parsed_meta["title"] == "新笔记"
    assert parsed_content.strip() == "正文内容"


def test_wikilink_and_tags() -> None:
    content = "参考 [[历史标书/2025项目]] 与 #智慧交通 #信号控制 的内容。"
    assert extract_wikilinks(content) == ["历史标书/2025项目"]
    tags = extract_tags(content, {"tags": ["企业资质"]})
    assert "企业资质" in tags
    assert "智慧交通" in tags
    assert "信号控制" in tags


def test_write_read_list_and_safe_path(tmp_path: Path) -> None:
    root = tmp_path / "kb"
    write_note(root, "01_企业资质/营业执照.md", "正文内容", {"title": "营业执照", "tags": ["资质"]})
    note = read_note(root, "01_企业资质/营业执照.md")
    assert note.metadata["title"] == "营业执照"
    assert "正文内容" in note.content
    assert list_notes(root) == ["01_企业资质/营业执照.md"]

    with pytest.raises(ValueError):
        safe_relative_path("../secret.md")
    with pytest.raises(ValueError):
        read_note(root, "../secret.md")


def test_tag_search(tmp_path: Path) -> None:
    root = tmp_path / "kb"
    write_note(root, "01_企业资质/ISO9001认证.md", "证书内容", {"tags": ["ISO9001"]})
    write_note(root, "02_历史标书/交通项目.md", "项目内容 #智慧交通", {})
    assert search_by_tag(root, "ISO9001") == ["01_企业资质/ISO9001认证.md"]
    assert search_by_tag(root, "智慧交通") == ["02_历史标书/交通项目.md"]


def test_gateway_initialize_creates_default_structure(tmp_path: Path) -> None:
    gateway = UnifiedKnowledgeGateway(_config(tmp_path))
    result = gateway.initialize()
    assert result["notes_count"] >= 1
    assert (tmp_path / "kb" / "README.md").exists()
    assert (tmp_path / "kb" / "01_企业资质" / ".gitkeep").exists()
    assert gateway.status()["pinakes_available"] is False


def test_keyword_search_fallback(tmp_path: Path) -> None:
    gateway = UnifiedKnowledgeGateway(_config(tmp_path))
    gateway.initialize()
    gateway.write_note("02_历史标书/智慧交通项目.md", "本项目为智慧交通信号控制系统建设。", {"tags": ["智慧交通"]})
    result = gateway.semantic_search("智慧交通 信号控制")
    assert result["engine"] == "keyword"
    assert result["results"]
    assert result["results"][0]["path"] == "02_历史标书/智慧交通项目.md"


def test_context_builder(tmp_path: Path) -> None:
    config = _config(tmp_path)
    root = config.root_path
    write_note(root, "01_企业资质/营业执照.md", "营业执照内容", {"tags": ["资质"]})
    write_note(root, "02_历史标书/智慧交通项目.md", "智慧交通信号控制系统技术方案。", {"industry": "IT"})
    write_note(root, "05_投标模板/IT/技术标模板.md", "# 技术标模板", {"industry": "IT"})
    write_note(root, "05_投标模板/通用/商务标模板.md", "# 商务标模板", {})
    write_note(root, "03_行业政策/国网招标规范.md", "国家电网招标规范要求。", {"industry": "IT"})
    write_note(root, "03_行业政策/政府采购法.md", "政府采购法律规定。", {"industry": "政务"})

    adapter = PinakesAdapter(config)
    builder = ContextBuilder(config, adapter)
    context = builder.get_context_for_bid("智慧交通信号控制系统", "IT")
    assert len(context["qualifications"]) == 1
    assert len(context["templates"]) == 2
    assert {item["path"] for item in context["templates"]} == {
        "05_投标模板/IT/技术标模板.md",
        "05_投标模板/通用/商务标模板.md",
    }
    assert len(context["policies"]) == 1
    assert context["policies"][0]["path"] == "03_行业政策/国网招标规范.md"


def test_keyword_extraction() -> None:
    keywords = extract_keywords("建设智慧交通信号控制系统与智能路侧系统", max_keywords=5)
    assert isinstance(keywords, list)
    assert all(len(word) >= 2 for word in keywords)


def test_reset_gateway_singleton(tmp_path: Path) -> None:
    from app.kb_integration import get_gateway

    config = _config(tmp_path)
    gateway = reset_gateway(config)
    assert get_gateway() is gateway
    assert gateway.root == tmp_path / "kb"


def test_chapter_mapper_places_content_into_correct_chapter(tmp_path: Path) -> None:
    config = _config(tmp_path)
    gateway = UnifiedKnowledgeGateway(config)
    gateway.initialize()
    gateway.write_note("01_企业资质/营业执照.md", "本公司营业执照内容。", {"tags": ["资质", "营业执照"]})
    gateway.write_note("04_公司介绍/公司简介.md", "本公司成立于2010年，主营信息通信与网络安全。", {"tags": ["公司介绍"]})
    gateway.write_note("02_历史标书/智慧交通项目.md", "智慧交通信号控制系统技术方案，采用分布式架构。", {"tags": ["成功案例", "技术方案"]})
    gateway.write_note("03_行业政策/国网招标规范.md", "国家电网招标规范要求。", {"tags": ["政策"]})

    mapper = ChapterKnowledgeMapper()
    chapters = [
        ("公司简介与项目业绩", []),
        ("资质证明", []),
        ("技术方案", []),
        ("合规性说明", []),
    ]
    knowledge = mapper.build_chapter_map(chapters, {}, "IT", gateway)
    assert any("公司简介" in item["path"] for item in knowledge.get("公司简介与项目业绩", []))
    assert any("营业执照" in item["path"] for item in knowledge.get("资质证明", []))
    assert any("智慧交通项目" in item["path"] for item in knowledge.get("技术方案", []))
    assert any("国网招标规范" in item["path"] for item in knowledge.get("合规性说明", []))
    assert format_knowledge_items(knowledge["资质证明"])  # 非空且可格式化


def test_chapter_mapper_wikilink_expansion_and_frontmatter_filter(tmp_path: Path) -> None:
    config = _config(tmp_path)
    gateway = UnifiedKnowledgeGateway(config)
    gateway.initialize()
    gateway.write_note(
        "01_企业资质/营业执照.md",
        "本公司营业执照内容。参考 [[历史业绩]]。",
        {"tags": ["资质"], "industry": "IT"},
    )
    gateway.write_note("02_历史标书/历史业绩.md", "近三年类似项目业绩清单。", {"tags": ["业绩"]})
    gateway.write_note("01_企业资质/IT资质.md", "IT 行业资质。", {"tags": ["资质"], "industry": "IT"})
    gateway.write_note("01_企业资质/政务资质.md", "政务行业资质。", {"tags": ["资质"], "industry": "政务"})

    rule = {
        "keywords": ["资质证明"],
        "directories": ["01_企业资质"],
        "tags": ["资质"],
        "frontmatter": ["industry:IT"],
        "query": "",
        "max_items": 5,
    }
    mapper = ChapterKnowledgeMapper([rule])
    items = mapper.retrieve_for_chapter("资质证明", {}, "IT", gateway)
    paths = [item["path"] for item in items]
    assert any("营业执照" in path for path in paths)
    assert any("历史业绩" in path for path in paths)
    assert any("IT资质" in path for path in paths)
    assert not any("政务资质" in path for path in paths)


def test_bid_special_rules_defaults_and_reset(tmp_path: Path, monkeypatch) -> None:
    from app.kb_integration import bid_rules_store

    monkeypatch.setattr(bid_rules_store, "_rules_path", tmp_path / "bid_chapter_rules.json")
    monkeypatch.setattr(bid_rules_store, "_cache", None)
    rules = bid_rules_store.get_bid_chapter_rules()
    assert len(rules) == len(bid_rules_store.DEFAULT_BID_CHAPTER_RULES)
    assert any("销售业绩" in rule["keywords"] for rule in rules)
    assert any("认证证书" in rule["keywords"] for rule in rules)

    custom = [
        {
            "keywords": ["销售业绩"],
            "directories": ["02_历史标书"],
            "tags": ["业绩"],
            "frontmatter": ["status:已中标"],
            "query": "类似项目业绩",
            "max_items": 4,
        }
    ]
    saved = bid_rules_store.save_bid_chapter_rules(custom)
    assert saved[0]["max_items"] == 4
    assert bid_rules_store.get_bid_chapter_rules() == saved

    reset = bid_rules_store.reset_bid_chapter_rules()
    assert reset == bid_rules_store.DEFAULT_BID_CHAPTER_RULES


def test_bid_mapper_uses_special_rules(tmp_path: Path) -> None:
    from app.kb_integration.bid_rules_store import get_bid_chapter_rules
    from app.kb_integration.chapter_mapper import ChapterKnowledgeMapper

    config = _config(tmp_path)
    gateway = UnifiedKnowledgeGateway(config)
    gateway.initialize()
    gateway.write_note(
        "02_历史标书/国网业绩.md",
        "近三年完成国网安全设备供货合同，合同金额860万元，已验收。",
        {"tags": ["业绩", "中标"], "status": "已中标"},
    )
    gateway.write_note(
        "02_历史标书/普通案例.md",
        "某园区网络建设项目。",
        {"tags": ["项目案例"]},
    )
    mapper = ChapterKnowledgeMapper(get_bid_chapter_rules())
    items = mapper.retrieve_for_chapter("销售业绩（对应招标公告资格条件要求）", {}, "IT", gateway)
    paths = [item["path"] for item in items]
    assert any("国网业绩" in path for path in paths)
    assert not any("普通案例" in path for path in paths)


def test_inject_bid_knowledge_into_doc(tmp_path: Path) -> None:
    from docx import Document

    from app.modules.bid_knowledge import inject_knowledge_into_doc

    doc = Document()
    doc.add_paragraph("销售业绩（对应招标公告资格条件要求）")
    knowledge = {
        "销售业绩（对应招标公告资格条件要求）": [
            {
                "path": "02_历史标书/国网业绩.md",
                "score": 1.0,
                "content": "近三年完成国网安全设备供货合同。",
                "metadata": {},
                "source": "目录",
            }
        ]
    }
    count = inject_knowledge_into_doc(doc, knowledge)
    assert count == 1
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "知识库参考内容" in text
    assert "近三年完成国网安全设备供货合同" in text
