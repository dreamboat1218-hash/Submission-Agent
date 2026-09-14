from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.database import get_session_factory
from app.models import Task
from app.modules import package_bid


def test_generate_package_bid_docx(tmp_path):
    md_path = tmp_path / "产品技术规范书.md"
    md_path.write_text(
        "1.技术参数特性表\n"
        "| 序号 | 名称 | 单位 | 项目需求值或表述 | 投标人保证值 |\n"
        "|---|---|---|---|---|\n"
        "| 1 | 部署要求 | 套 | 旁路部署 |  |\n"
        "2.组件材料配置表\n"
        "| 序号 | 组件材料名称 | 规格型号 | 单位 | 数量 | 备注 |\n"
        "|---|---|---|---|---|---|\n"
        "| 1 | 加密引擎 | XX | 个 | 2 |  |\n",
        encoding="utf-8",
    )
    packages_data = {
        "project": {"项目名称": "测试项目", "项目编号": "TEST-001", "分标名称": "测试分标", "分标编号": "TEST-1"},
        "packages": [
            {
                "package_no": "包1",
                "包号": "包1",
                "包名称": "包1测试设备",
                "files": [
                    {
                        "name": "产品技术规范书.docx",
                        "md_path": str(md_path),
                        "kind": "tech_spec",
                        "has_tech_params": True,
                        "tech_spec_code": "B005-500000000-00001",
                    }
                ],
            }
        ],
    }

    factory = get_session_factory()
    with factory() as db:
        task = Task(user_id="test", status="completed", original_file_name="tender.rar", work_dir=str(tmp_path))
        db.add(task)
        db.commit()
        db.refresh(task)
        task_id = task.id

    try:
        result = package_bid.generate_package_bid(task_id, "包1", [str(md_path)], packages_data)
        assert result["product_count"] == 1
        assert result["products"] == ["产品技术规范书.docx"]

        doc = Document(result["bid_file_path"])
        assert len(doc.tables) == 3
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "项目货物技术特性逐条逐项应答文件" in text
        assert "项目货物组件材料配置表" in text
        assert "技术偏差表" in text
        assert doc.tables[0].cell(1, 0).text == "1"
        assert doc.tables[1].cell(1, 1).text == "加密引擎"
    finally:
        with factory() as db:
            task = db.get(Task, task_id)
            if task is not None:
                db.delete(task)
                db.commit()


def test_restyle_technical_heads_levels() -> None:
    doc = Document()

    def add(text: str, outline: int | None = None):
        paragraph = doc.add_paragraph(text)
        if outline is not None:
            p_pr = paragraph._p.get_or_add_pPr()
            level = OxmlElement("w:outlineLvl")
            level.set(qn("w:val"), str(outline))
            p_pr.append(level)
        return paragraph

    title = add("三、技术文件", 1)
    header = add("（一）技术偏差表", 2)
    numbered = add("1.技术偏差表", 3)
    table_title = add("技术偏差表")

    package_bid._restyle_technical_heads(doc)

    assert title.style.name == "Normal"
    assert title.text == "三、技术文件"
    assert header.style.name == "Normal"
    assert header.text == "（一）技术偏差表"
    assert numbered.style.name == "Heading 1"
    assert numbered.text == "技术偏差表"
    assert table_title.style.name == "Normal"


def test_restart_section_numbering_uses_new_num_id_per_section() -> None:
    template = Path(__file__).resolve().parents[1] / "templates" / "投标格式.docx"
    doc = Document(str(template))
    doc.add_paragraph("（一）技术偏差表")
    first = doc.add_paragraph("技术偏差表")
    first.style = doc.styles["Heading 1"]
    doc.add_paragraph("（二）专项投标文件")
    second = doc.add_paragraph("业绩文件")
    second.style = doc.styles["Heading 1"]

    package_bid._restart_section_numbering(doc)

    first_id = first._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val
    second_id = second._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val
    assert first_id is not None
    assert second_id is not None
    assert first_id != second_id


def test_drop_empty_heading_paragraphs() -> None:
    doc = Document()
    empty_heading = doc.add_paragraph()
    empty_heading.style = doc.styles["Heading 2"]
    empty_heading_9 = doc.add_paragraph("   ")
    empty_heading_9.style = doc.styles["Heading 9"]
    keep = doc.add_paragraph("技术偏差表")
    keep.style = doc.styles["Heading 1"]
    blank_normal = doc.add_paragraph("")

    package_bid._drop_empty_heading_paragraphs(doc)

    assert [paragraph.text for paragraph in doc.paragraphs] == ["技术偏差表", ""]
