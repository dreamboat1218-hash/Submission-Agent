from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

from docx import Document

from app.modules import fallback_extractor
from app.modules.generator import _build_chapter_plan, _build_raw_from_source_doc, _build_raw_template_doc

SAMPLE = Path(__file__).parent / "sample_tender.md"


def test_dynamic_chapter_plan_follows_bid_format():
    content = SAMPLE.read_text(encoding="utf-8")
    data = fallback_extractor.extract_file(content, "招标公告.md")
    chapters = _build_chapter_plan(data, "IT", "知识库上下文")

    titles = [title for title, _ in chapters]
    assert "商务文件" in titles
    assert "技术文件" in titles
    assert "价格文件" in titles
    assert "投标函" in titles
    assert "响应性说明与格式符合性" in titles


def test_dynamic_docx_contains_tender_sections(tmp_path):
    content = SAMPLE.read_text(encoding="utf-8")
    data = fallback_extractor.extract_file(content, "招标公告.md")
    chapters = _build_chapter_plan(data, "IT", "")

    doc = Document()
    for index, (title, sections) in enumerate(chapters, start=1):
        doc.add_paragraph(f"第{index}章 {title}")
        for section_title, body in sections:
            doc.add_paragraph(section_title)
            doc.add_paragraph(body[:50])
    path = tmp_path / "dynamic.docx"
    doc.save(path)

    text = "\n".join(p.text for p in Document(path).paragraphs)
    assert "商务文件" in text
    assert "价格文件" in text
    assert "投标函" in text


def test_raw_template_doc_only_contains_commercial_and_technical():
    content = SAMPLE.read_text(encoding="utf-8")
    data = fallback_extractor.extract_file(content, "招标公告.md")

    doc = Document()
    assert _build_raw_template_doc(doc, data, "IT") is True
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "商务文件" in text
    assert "技术文件" in text
    assert "商务偏差表" in text
    assert "技术偏差表" in text
    assert "价格文件" not in text


def test_source_docx_copy_preserves_tables_and_layout(tmp_path):
    source = tmp_path / "tender_format.docx"
    doc = Document()
    doc.add_paragraph("第六章  投标文件格式")
    doc.add_paragraph("商务部分（单独成册）")
    table1 = doc.add_table(rows=2, cols=3)
    table1.style = "Table Grid"
    table1.cell(0, 0).text = "序号"
    table1.cell(0, 1).text = "招标文件条款"
    table1.cell(0, 2).text = "投标文件条款"
    table1.cell(1, 0).merge(table1.cell(1, 1))
    table1.cell(1, 0).text = "合并单元格"
    table1.cell(1, 2).text = ""
    doc.add_paragraph("商务偏差表说明")
    doc.add_paragraph("技术部分（单独成册）")
    table2 = doc.add_table(rows=2, cols=2)
    table2.style = "Table Grid"
    table2.cell(0, 0).text = "偏差事项"
    table2.cell(0, 1).text = "招标文件要求"
    table2.cell(1, 0).text = "1"
    table2.cell(1, 1).text = "2"
    doc.add_paragraph("技术偏差表说明")
    doc.add_paragraph("投标保证保险（责任要求）")
    doc.add_paragraph("不应保留的内容")
    doc.save(source)

    data = {
        "project_name": {"primary_value": "测试项目", "confidence": 0.9, "evidence": {}},
        "bid_number": {"primary_value": "TEST-001", "confidence": 0.9, "evidence": {}},
        "bid_format": {
            "primary_value": {
                "commercial_raw": "商务部分（单独成册）",
                "technical_raw": "技术部分（单独成册）",
                "source_path": str(source),
            },
            "confidence": 0.9,
            "evidence": {},
        },
    }
    target = tmp_path / "output.docx"
    result = _build_raw_from_source_doc(data, "IT", source, tmp_path, target)
    assert result == target

    output = Document(target)
    text = "\n".join(p.text for p in output.paragraphs)
    assert "投 标 文 件" in text
    assert "商务部分（单独成册）" in text
    assert "技术部分（单独成册）" in text
    assert "第六章  投标文件格式" not in text
    assert "不应保留的内容" not in text
    assert len(output.tables) == 2
    assert output.tables[0].cell(1, 0).text == "合并单元格"


def test_kb_attachment_injected_into_corresponding_chapter(tmp_path):
    source = tmp_path / "tender_format.docx"
    doc = Document()
    doc.add_paragraph("商务部分（单独成册）")
    doc.add_paragraph("商务内容")
    doc.add_paragraph("技术部分（单独成册）")
    doc.add_paragraph("技术内容")
    doc.save(source)

    data = {
        "project_name": {"primary_value": "测试项目", "confidence": 0.9, "evidence": {}},
        "bid_number": {"primary_value": "TEST-001", "confidence": 0.9, "evidence": {}},
        "bid_format": {
            "primary_value": {"commercial_raw": "商务", "technical_raw": "技术", "source_path": str(source)},
            "confidence": 0.9,
            "evidence": {},
        },
    }
    kb_docs = [
        SimpleNamespace(title="企业营业执照附件", content="营业执照扫描件内容", metadata_json={"chapter": "商务文件", "images": []}),
        SimpleNamespace(title="技术方案附件", content="微服务架构技术方案", metadata_json={"chapter": "技术文件", "images": []}),
    ]
    target = tmp_path / "output.docx"
    _build_raw_from_source_doc(data, "IT", source, tmp_path, target, kb_docs)

    text = "\n".join(p.text for p in Document(target).paragraphs)
    assert "企业营业执照附件" in text
    assert "技术方案附件" in text
    commercial_index = text.index("商务部分（单独成册）")
    technical_index = text.index("技术部分（单独成册）")
    assert commercial_index < text.index("企业营业执照附件") < technical_index
    assert technical_index < text.index("技术方案附件")
